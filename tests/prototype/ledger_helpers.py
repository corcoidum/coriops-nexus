"""시험용 합성 장부 엑셀과 고정 시계. 실제 환자 값은 쓰지 않는다."""

from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO

from openpyxl import Workbook

# 2026-11-05(목) 10:00 KST
THURSDAY_10 = datetime(2026, 11, 5, 1, 0, tzinfo=timezone.utc)
SENTINEL = "SENTINEL-비고-값"

BIOPSY_HEADERS = ["검사일", "접수번호", "차트번호", "수진자명", "나이", "성별", "위탁일", "접수인", "검사종류",
                  "검사위치", "검체수량", "인계인", "인수인", "결과보고일", "조직검사결과", "특이소견", "비고란"]
SHORT_HEADERS = ["검사일", "접수번호", "차트번호", "수진자명", "나이", "성별", "위탁일", "결과보고일", "검사결과", "비고"]


class FixedClock:
    def __init__(self, moment: datetime = THURSDAY_10):
        self.moment = moment

    def now(self) -> datetime:
        return self.moment


def biopsy_row(day: date, ref: str, chart: str, name: str, *, kind="stomach", result="H#1", reported=None,
               sydney="", entrusted=True, count=1, note=SENTINEL):
    return [day, ref, chart, name, 50, "M", day if entrusted else None, "테스트 접수자", kind, "합성 위치", count,
            "테스트 인계자", "테스트 수령인", reported, result, sydney, note]


def short_row(day: date, ref: str, chart: str, name: str, result="-", reported=None):
    return [day, ref, chart, name, 60, "F", day, reported, result, SENTINEL]


def workbook_bytes(biopsy=(), clo=(), stool=(), extra_sheet=True, biopsy_headers=BIOPSY_HEADERS) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Biopsy"
    sheet.append(biopsy_headers)
    for row in biopsy:
        sheet.append(row)
    for title, rows in (("CLO", clo), ("Stool", stool)):
        other = workbook.create_sheet(title)
        other.append(SHORT_HEADERS)
        for row in rows:
            other.append(row)
    if extra_sheet:
        workbook.create_sheet("분석 대시보드").append(["합성 통계"])
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


def standard_workbook() -> bytes:
    """과거 4건 + 오늘 1건(인계 대기) + CLO 1건 + Stool 1건."""
    return workbook_bytes(
        biopsy=[
            biopsy_row(date(2026, 10, 2), "T9800001", "H-0001", "가져온환자일", result="TA#2(L), H#1",
                       reported=date(2026, 10, 5), kind="colon", count=2),
            biopsy_row(date(2026, 10, 6), "T9800002", "H-0002", "가져온환자이", result="CAG#1(im)",
                       reported=date(2026, 10, 7), sydney="HP-/Neu++/Mon++/Atr+/IM-"),
            biopsy_row(date(2026, 10, 8), "T9800003", "H-0003", "가져온환자삼", result="tubular adenoma 문장",
                       reported=date(2026, 10, 10)),
            biopsy_row(date(2026, 10, 12), "", "H-0004", "가져온환자사", result="", entrusted=False),
            biopsy_row(date(2026, 11, 5), "", "H-0005", "가져온환자오", result="", entrusted=False),
        ],
        clo=[short_row(date(2026, 10, 20), "", "H-0002", "가져온환자이", result="+", reported=date(2026, 10, 21))],
        stool=[short_row(date(2026, 10, 21), "T9900001", "H-0006", "가져온환자육", result="-", reported=date(2026, 10, 23))],
    )
