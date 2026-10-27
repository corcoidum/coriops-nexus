"""대장별 분리 관리(A18): CLO 대장 5열(검사일·차트번호·수진자명·성별·CLO결과) 가져오기·출력, 대장별 메뉴."""

from __future__ import annotations

from datetime import date
from io import BytesIO
import html
import re

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from app.main import create_app
from app.modules.endoscopy_pathology.excel_io import LedgerAppConfig
from app.modules.endoscopy_pathology.ledger_db import MemoryStore
from app.modules.endoscopy_pathology.manual_dashboard import DashboardError, ManualDashboardService
from ledger_helpers import FixedClock, standard_workbook

CLO_HEADERS = ["검사일", "차트번호", "수진자명", "성별", "CLO결과"]


def clo_file(*sheets: tuple[str, list]) -> bytes:
    """원내 수기 CLO 대장과 같은 5열. 시트 이름은 정해 두지 않는다."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        if rows is not None:
            sheet.append(CLO_HEADERS)
            for row in rows:
                sheet.append(row)
        else:
            sheet.append(["합성 요약"])
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


ROWS_2025 = [[date(2025, 12, 1), "C-0001", "테스트씨엘오일", "M", "+"],
             [date(2025, 12, 2), "C-0002", "테스트씨엘오이", "F", "-"]]
ROWS_2026 = [[date(2026, 10, 30), "C-0003", "테스트씨엘오삼", "남", "양성"],
             [date(2026, 11, 4), "C-0004", "테스트씨엘오사", "F", None]]


@pytest.fixture
def service(tmp_path):
    return ManualDashboardService(clock=FixedClock(), store=MemoryStore(), export_dir=tmp_path)


def test_clo_file_is_read_as_clo_whatever_the_sheet_names(service):
    token = service.preview_import(clo_file(("2025", ROWS_2025), ("2026", ROWS_2026), ("요약", None)), kind="CLO")
    preview = service.preview(token)
    assert preview["new_total"] == 4
    assert service.commit_import(token) == 4
    entries = {e.chart_no: e for e in service._store.entries()}
    assert {e.kind for e in entries.values()} == {"CLO"}
    # 결과가 적힌 행은 결과 날짜 = 검사일, 결과 기록으로 종결
    first = entries["C-0001"]
    assert (first.result_code, first.reported_on, first.sex) == ("+", date(2025, 12, 1), "M")
    assert entries["C-0003"].sex == "M"
    stages = service.snapshot()["stages"]
    assert [stages[entries[c].id] for c in ("C-0001", "C-0002", "C-0003")] == ["CLOSED"] * 3
    # 결과가 비어 있으면 결과 대기로 남는다(추측하지 않음)
    assert (entries["C-0004"].reported_on, stages[entries["C-0004"].id]) == (None, "RESULT_PENDING")


def test_clo_result_column_is_recognised_in_the_automatic_mode(service):
    token = service.preview_import(clo_file(("CLO", ROWS_2025)))
    preview = service.preview(token)
    assert preview["new_total"] == 2
    assert [s["unknown_headers"] for s in preview["sheets"]] == [[]]  # CLO결과·성별 모두 인식
    service.commit_import(token)
    assert {e.result_code for e in service._store.entries()} == {"+", "-"}


def test_unknown_sheet_names_need_the_ledger_choice(service):
    with pytest.raises(DashboardError) as error:
        service.preview_import(clo_file(("2025", ROWS_2025)))
    assert "대장 종류" in str(error.value)
    with pytest.raises(DashboardError) as no_headers:
        service.preview_import(clo_file(("요약", None)), kind="CLO")
    assert "필수 열" in str(no_headers.value)
    with pytest.raises(DashboardError):
        service.preview_import(clo_file(("2025", ROWS_2025)), kind="LAB")


def test_each_ledger_exports_alone_and_clo_has_exactly_five_columns(service, tmp_path):
    service.commit_import(service.preview_import(standard_workbook()))
    service.commit_import(service.preview_import(clo_file(("2025", ROWS_2025)), kind="CLO"))

    record = service.export("CLO")
    assert record.file_name.startswith("CLO대장_")
    workbook = load_workbook(tmp_path / record.file_name)
    assert workbook.sheetnames == ["CLO"]
    sheet = workbook.active
    assert [c.value for c in sheet[1]] == CLO_HEADERS
    values = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
    assert ["C-0001", "테스트씨엘오일", "M", "+"] in [row[1:] for row in values]
    assert sheet.max_column == 5  # 재검 열 없이 수기 대장 그대로

    biopsy = load_workbook(tmp_path / service.export("BIOPSY").file_name)
    assert biopsy.sheetnames == ["Biopsy"]
    everything = load_workbook(tmp_path / service.export().file_name)
    assert everything.sheetnames == ["Biopsy", "CLO", "Stool"]
    assert [c.value for c in everything["CLO"][1]] == CLO_HEADERS
    with pytest.raises(DashboardError):
        service.export("LAB")


def test_export_headers_must_be_known_columns(tmp_path):
    source = (LedgerAppConfig.load.__func__.__globals__["DEFAULT_CONFIG"]).read_text(encoding="utf-8")
    broken = tmp_path / "ledger_app.yaml"
    broken.write_text(source.replace("export_headers: [검사일, 차트번호, 수진자명, 성별, CLO결과]",
                                     "export_headers: [검사일, 비고]"), encoding="utf-8")
    with pytest.raises(ValueError):
        LedgerAppConfig.load(broken)


# ---------- 화면: 대장별 메뉴


@pytest.fixture
def client(tmp_path):
    service = ManualDashboardService(clock=FixedClock(), export_dir=tmp_path)  # 고정 합성 14건
    app = create_app(service)
    with TestClient(app) as test_client:
        test_client.app_state = app.state
        yield test_client


def test_ledger_menus_and_clo_screen(client):
    home = client.get("/dashboard").text
    for label, kind in (("조직검체 대장", "BIOPSY"), ("CLO 대장", "CLO"), ("분변 대장", "STOOL")):
        assert f'href="/dashboard?kind={kind}"' in home and label in home
    clo = client.get("/dashboard?kind=CLO").text
    assert "<h1>CLO 대장</h1>" in clo
    for header in CLO_HEADERS:
        assert f"<th scope=\"col\">{header}</th>" in clo
    assert "오늘 수탁기관 인계" not in clo  # CLO는 인계가 없다
    assert "원장 지시 대기" not in clo.split("work-grid")[0]  # 단계 카드에 없음
    assert "CLO 대장 엑셀 저장" in clo
    links = {html.unescape(h) for h in re.findall(r'href="(/dashboard[^"]*)"', clo)}
    for href in links:
        assert client.get(href).status_code == 200, href


def test_ledger_export_from_the_ledger_page_returns_there(client):
    response = client.post("/ledger/export", data={"csrf_token": client.app_state.csrf_token, "kind": "CLO",
                                                   "back": "ledger"}, follow_redirects=False)
    assert response.status_code == 303
    assert response.headers["location"].startswith("/dashboard?kind=CLO")
    assert "notice=exported" in response.headers["location"]
    assert "지정 폴더에 엑셀을 저장했습니다." in client.get(response.headers["location"]).text


def test_import_form_offers_the_ledger_choice(client):
    page = client.get("/ledger/excel").text
    assert 'name="ledger_kind"' in page and "CLO 대장 파일" in page
