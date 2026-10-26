"""재검 도래·기간 필터·대상자 목록 출력. 시계는 2026-11-05(목) 10:00 KST로 고정한다."""

from __future__ import annotations

from datetime import date

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.main import create_app
from app.modules.endoscopy_pathology.manual_dashboard import DashboardError, ManualDashboardService
from ledger_helpers import FixedClock

TODAY = date(2026, 11, 5)


def followup(service, chart, due, kind="BIOPSY"):
    """합성 재검 예정 1건: 접수 → 인계 → 결과 → 원장 지시."""
    extra = {"procedure_type": "colon", "specimen_count": "1"} if kind == "BIOPSY" else {}
    entry_id = service.register(kind, chart, f"재검{chart[-2:]}", date(2025, 11, 5), **extra)
    pending = [e for e in service._store.entries() if service._stage(e) == "HANDOFF_PENDING"]
    if pending:
        service.handoff("a", "b", "c", sum(e.specimen_count or 1 for e in pending))
    service.receive(entry_id, f"R{chart}", "TA#1(L)" if kind == "BIOPSY" else "+", "합성 위치")
    service.decide(entry_id, "required", due)
    return entry_id


@pytest.fixture
def service(tmp_path):
    service = ManualDashboardService(clock=FixedClock(), export_dir=tmp_path)
    # seed: 10/26(경과), 11/25(D-20), 2027-09-01(D-300) + 오늘·이번 주·다음 달 추가
    followup(service, "F-0001", TODAY)
    followup(service, "F-0002", date(2026, 11, 8))      # 이번 주 일요일
    followup(service, "F-0003", date(2026, 12, 5))      # 30일째
    followup(service, "F-0004", date(2026, 11, 2), kind="CLO")  # 이번 주 월요일, 경과
    return service


def charts(snapshot):
    return [e.chart_no for e in snapshot["cases"]]


def test_due_today_is_flagged_and_counted_as_arrived(service):
    snap = service.snapshot()
    today_entry = next(e for e in service._store.entries() if e.chart_no == "F-0001")
    alert = snap["alerts"][today_entry.id]

    assert (alert.text, alert.level) == ("재검일 오늘", "danger")
    assert snap["flags"]["followup_due"] == 3     # 오늘 + 경과 2건(10/26, 11/02)
    assert snap["flags"]["followup_overdue"] == 2


@pytest.mark.parametrize(("due", "expected"), [
    ("due", ["T-0010", "F-0004", "F-0001"]),
    ("week", ["F-0004", "F-0001", "F-0002"]),
    ("month", ["F-0004", "F-0001", "F-0002", "T-0009"]),
    ("30d", ["F-0001", "F-0002", "T-0009", "F-0003"]),
    ("", None),
])
def test_due_filters_select_by_followup_date_in_date_order(service, due, expected):
    snap = service.snapshot(due=due, stage="FOLLOWUP")
    if expected is None:
        assert snap["matched"] == snap["counts"]["FOLLOWUP"] == 7
    else:
        assert charts(snap) == expected
        assert snap["stage"] == "FOLLOWUP"


def test_custom_range_and_its_validation(service):
    snap = service.snapshot(due="range", due_from=date(2026, 11, 8), due_to=date(2026, 12, 5))
    assert charts(snap) == ["F-0002", "T-0009", "F-0003"]
    for start, end in [(None, date(2026, 12, 1)), (date(2026, 12, 2), date(2026, 12, 1)),
                       (date(2026, 1, 1), date(2028, 6, 1))]:
        with pytest.raises(DashboardError) as error:
            service.snapshot(due="range", due_from=start, due_to=end)
        assert error.value.status_code == 422
    with pytest.raises(DashboardError):
        service.snapshot(due="tomorrow")


def test_due_filter_combines_with_kind_and_search(service):
    assert charts(service.snapshot(due="due", kind="CLO")) == ["F-0004"]
    assert charts(service.snapshot(due="due", query="F-0001")) == ["F-0001"]


def test_due_filter_page_menu_badge_and_links(service):
    app = create_app(service)
    with TestClient(app) as client:
        page = client.get("/dashboard?stage=FOLLOWUP&due=due").text
        bad = client.get("/dashboard?due=range&due_from=2026-13-01&due_to=2026-12-01")
    assert "재검 대상 · 재검 도래" in page
    assert 'class="nav-badge">3<' in page
    assert "kind=CLO&amp;due=due" in page  # 종류 탭으로 바꿔도 도래 조건 유지
    assert "이 목록 엑셀 저장 (3명)" in page
    assert bad.status_code == 422


def test_followup_list_export_writes_only_the_filtered_people(service, tmp_path):
    record, subjects = service.export_followups(due="due")

    assert record.file_name == "재검대상_20261105_100000.xlsx" and record.row_count == 3
    assert subjects == ["T-0010", "F-0004", "F-0001"]
    sheet = load_workbook(tmp_path / record.file_name)["재검 대상"]
    assert sheet["A1"].value.startswith("재검 대상자 · 재검 도래 · 처음 ~ 2026-11-05 · 조직검체 장부 · 3명")
    assert [c.value for c in sheet[3]][:7] == ["재검 예정일", "상태", "장부", "검사종류", "검사일", "차트번호", "수진자명"]
    rows = list(sheet.iter_rows(min_row=4, values_only=True))
    assert [(r[1], r[5]) for r in rows] == [("재검 기한 10일 경과", "T-0010"), ("재검 기한 3일 경과", "F-0004"),
                                            ("재검일 오늘", "F-0001")]
    assert rows[1][2] == "RUT(CLO)" and rows[2][6] == "재검01"  # 연락용 목록이라 이름은 그대로
    assert service.snapshot()["exports"][0].file_name == record.file_name


def test_followup_export_refuses_empty_lists_and_missing_folder(service, tmp_path):
    with pytest.raises(DashboardError) as empty:
        service.export_followups(kind="STOOL", due="due", ledger="fobt")
    assert empty.value.status_code == 422
    no_folder = ManualDashboardService(clock=FixedClock(), export_dir="")
    with pytest.raises(DashboardError) as missing:
        no_folder.export_followups(due="due")
    assert missing.value.status_code == 409
    assert list(tmp_path.iterdir()) == []


def test_followup_export_route_keeps_the_filter_and_needs_csrf(service, tmp_path):
    app = create_app(service)
    form = {"kind": "ALL", "q": "", "due": "week", "due_from": "", "due_to": ""}
    with TestClient(app) as client:
        denied = client.post("/ledger/export/followups", data={"csrf_token": "INVALID", **form})
        done = client.post("/ledger/export/followups", data={"csrf_token": app.state.csrf_token, **form},
                           follow_redirects=False)
    assert denied.status_code == 403
    assert done.status_code == 303 and "due=week" in done.headers["location"]
    files = list(tmp_path.iterdir())
    assert len(files) == 1 and load_workbook(files[0])["재검 대상"].max_row == 3 + 3
