"""CLO 대장 분리(A27): 5열(검사일·차트번호·수진자명·성별·CLO결과) 가져오기·출력, 결과 기록으로 종결, /clo 화면."""

from __future__ import annotations

from datetime import date
from io import BytesIO
import html
import re

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from app.main import create_app
from app.modules.endoscopy_pathology.excel_io import DEFAULT_CONFIG, LedgerAppConfig
from app.modules.endoscopy_pathology.ledger_db import MemoryStore
from app.modules.endoscopy_pathology.manual_dashboard import DashboardError, ManualDashboardService
from ledger_helpers import FixedClock, standard_workbook

CLO_HEADERS = ["검사일", "차트번호", "수진자명", "성별", "CLO결과"]


def clo_file(*sheets: tuple[str, list | None]) -> bytes:
    """원내 수기 CLO 대장과 같은 5열. 시트 이름은 정해 두지 않는다. rows가 None이면 요약표 시트."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        if rows is None:
            sheet.append(["합성 요약"])
            continue
        sheet.append(CLO_HEADERS)
        for row in rows:
            sheet.append(row)
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


ROWS_2025 = [[date(2025, 12, 1), "C-0001", "테스트씨엘오일", "M", "+"],
             [date(2025, 12, 2), "C-0002", "테스트씨엘오이", "F", "-"]]
ROWS_2026 = [[date(2026, 10, 30), "C-0003", "테스트씨엘오삼", "남", "양성"],
             [date(2026, 11, 4), "C-0004", "테스트씨엘오사", "F", None]]


@pytest.fixture
def service(tmp_path):
    return ManualDashboardService(clock=FixedClock(), store=MemoryStore(), export_dir=tmp_path)


def test_clo_file_is_read_as_clo_whatever_the_sheet_names(service):
    token = service.preview_import(clo_file(("2025", ROWS_2025), ("2026", ROWS_2026), ("요약", None)), ledger="clo")
    preview = service.preview(token)
    assert preview["new_total"] == 4 and preview["other_sheets"] == ["요약"]
    assert service.commit_import(token) == 4
    entries = {e.chart_no: e for e in service._store.entries()}
    assert {e.kind for e in entries.values()} == {"CLO"}
    # 회신날짜 열이 없으므로 결과가 적힌 행의 결과 날짜는 검사일, 결과 기록으로 종결
    first = entries["C-0001"]
    assert (first.result_code, first.reported_on, first.sex) == ("+", date(2025, 12, 1), "M")
    assert entries["C-0003"].sex == "M"
    view = service.snapshot(ledger="clo")
    assert [view["stages"][entries[c].id] for c in ("C-0001", "C-0002", "C-0003")] == ["CLOSED"] * 3
    # 결과가 비어 있으면 결과 대기로 남는다(추측하지 않음)
    assert (entries["C-0004"].reported_on, view["stages"][entries["C-0004"].id]) == (None, "RESULT_PENDING")
    assert view["total"] == 4 and service.snapshot()["total"] == 0  # 조직검체 장부에는 없다


def test_clo_result_column_is_recognised_in_the_automatic_mode(service):
    token = service.preview_import(clo_file(("CLO", ROWS_2025)))
    preview = service.preview(token)
    assert preview["new_total"] == 2
    assert [s["unknown_headers"] for s in preview["sheets"]] == [[]]  # CLO결과·성별 모두 인식
    service.commit_import(token)
    assert {e.result_code for e in service._store.entries()} == {"+", "-"}


def test_unknown_sheet_names_need_the_ledger_choice(service):
    with pytest.raises(DashboardError) as error:
        service.preview_import(clo_file(("2025", ROWS_2025)))
    assert "장부 종류" in str(error.value)
    with pytest.raises(DashboardError) as no_headers:
        service.preview_import(clo_file(("요약", None)), ledger="clo")
    assert "필수 열" in str(no_headers.value)
    with pytest.raises(DashboardError):
        service.preview_import(clo_file(("2025", ROWS_2025)), ledger="lab")


def test_clo_result_closes_without_a_doctor_step(service):
    entry_id = service.register("CLO", "C-0101", "테스트씨엘오", date(2026, 11, 5), sex="F")
    view = service.snapshot(ledger="clo")
    assert view["stages"][entry_id] == "RESULT_PENDING"
    assert view["stage_cards"] == ["RESULT_PENDING", "CLOSED"]  # 인계·원장 지시·재검 카드 없음
    assert not view["has_handoff"] and view["delay_days"] == 1  # RUT 결과 지연 기준
    service.receive(entry_id, "", "-", "")
    view = service.snapshot(ledger="clo")
    assert view["stages"][entry_id] == "CLOSED"
    assert view["alerts"][entry_id].text == "결과 기록 · 종결"
    with pytest.raises(DashboardError):
        service.decide(entry_id, "not_required", reason="합성")


def test_closed_clo_counts_for_retention_from_the_result_date(service):
    service.commit_import(service.preview_import(clo_file(("2025", ROWS_2025)), ledger="clo"))
    service._clock.moment = service._clock.moment.replace(year=2031)
    assert service.snapshot(ledger="clo")["retention"]["expired"] == 2


def test_clo_export_has_exactly_the_five_columns(service, tmp_path):
    service.commit_import(service.preview_import(standard_workbook()))
    service.commit_import(service.preview_import(clo_file(("2025", ROWS_2025)), ledger="clo"))

    record = service.export("clo")
    assert record.file_name.startswith("CLO대장_") and record.row_count == 3
    workbook = load_workbook(tmp_path / record.file_name)
    assert workbook.sheetnames == ["CLO"]
    sheet = workbook.active
    assert [c.value for c in sheet[1]] == CLO_HEADERS and sheet.max_column == 5  # 재검 열 없이 수기 대장 그대로
    rows = [[c.value for c in row][1:] for row in sheet.iter_rows(min_row=2)]
    assert ["C-0001", "테스트씨엘오일", "M", "+"] in rows

    assert load_workbook(tmp_path / service.export().file_name).sheetnames == ["Biopsy"]
    with pytest.raises(DashboardError):
        service.export("fobt")  # 분변잠혈은 종이 대장 ③ 양식으로 출력한다(A22)
    with pytest.raises(DashboardError):
        service.export("lab")


def test_export_headers_must_be_known_columns(tmp_path):
    broken = tmp_path / "ledger_app.yaml"
    broken.write_text(DEFAULT_CONFIG.read_text(encoding="utf-8").replace(
        "export_headers: [검사일, 차트번호, 수진자명, 성별, CLO결과]", "export_headers: [검사일, 비고]"), encoding="utf-8")
    with pytest.raises(ValueError):
        LedgerAppConfig.load(broken)


# ---------- 화면


@pytest.fixture
def client(tmp_path):
    service = ManualDashboardService(clock=FixedClock(), export_dir=tmp_path)  # 고정 합성 14건
    app = create_app(service)
    with TestClient(app) as test_client:
        test_client.app_state = app.state
        yield test_client


def post(client, path, **data):
    return client.post(path, data={"csrf_token": client.app_state.csrf_token, **data}, follow_redirects=False)


def test_clo_menu_and_screen(client):
    home = client.get("/").text
    assert 'href="/clo"' in home and "CLO 대장" in home
    page = client.get("/clo").text
    assert "<h1>CLO 대장</h1>" in page and "T-0013" in page and "T-0001" not in page
    for header in CLO_HEADERS:
        assert f'<th scope="col">{header}</th>' in page
    assert "오늘 수탁기관 인계" not in page  # CLO는 인계가 없다
    assert "원장 지시 대기" not in page.split("work-grid")[0]  # 단계 카드에 없음
    assert "CLO 대장 엑셀 저장" in page
    assert "테*****3" in page and "테스트환자13" in page  # 목록은 가리고, 선택한 상세는 전체 이름
    assert "검사 → 결과 기록 → 종결" in page.split("page-footer")[1]
    links = {html.unescape(h) for h in re.findall(r'href="(/clo[^"]*)"', page)}
    assert any("selected_id=" in h for h in links) and any("new=1" in h for h in links)
    for href in links:
        assert client.get(href).status_code == 200, href


def test_clo_register_and_result_stay_on_the_clo_ledger(client):
    form = client.get("/clo?new=1").text
    assert 'name="kind" value="CLO"' in form and "RUT(CLO) 1건" in form
    created = post(client, "/ledger/register", kind="CLO", exam_date="2026-11-05", chart_no="C-0201",
                   patient_name="테스트씨엘오", sex_age="F/61")
    location = created.headers["location"]
    assert created.status_code == 303 and location.startswith("/clo?")
    entry_id = location.split("selected_id=")[1].split("&")[0]
    result = post(client, f"/ledger/cases/{entry_id}/receive", result_notation="+")
    assert result.headers["location"] == f"/clo?selected_id={entry_id}"
    detail = client.get(result.headers["location"]).text
    assert "결과 기록 · 종결" in detail and "EMR에서" in detail and "원장 지시 기록" not in detail


def test_ledger_export_from_the_clo_page_returns_there(client):
    response = post(client, "/ledger/export", ledger="clo", back="ledger")
    assert response.status_code == 303 and response.headers["location"] == "/clo?notice=exported"
    assert "지정 폴더에 엑셀을 저장했습니다." in client.get(response.headers["location"]).text


def test_excel_page_offers_the_ledger_choice(client):
    page = client.get("/ledger/excel").text
    assert 'name="ledger"' in page and "CLO 대장 파일" in page
    assert "조직검체 장부 엑셀 저장" in page and "CLO 대장 엑셀 저장" in page
    assert "분변잠혈 대장 엑셀 저장" not in page
