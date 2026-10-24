"""합성 데이터 전용 장부 변환 계약: 실데이터·외부 조회를 사용하지 않는다."""

from __future__ import annotations

from dataclasses import replace
from datetime import date
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.modules.endoscopy_pathology.ledger import (
    ResultRecord,
    load_policy,
    merge_records,
    new_workbook,
    save_new_workbook,
)


SHEET_NAME = "조직검체장부"
HEADERS = (
    "검사일",
    "접수번호",
    "차트번호",
    "수진자명",
    "나이",
    "성별",
    "위탁일",
    "접수인",
    "검사종류",
    "검사위치",
    "검체수량",
    "인계인",
    "인수인",
    "결과보고일",
    "조직검사결과",
    "특이소견",
    "비고란",
)
MANUAL_HEADERS = ("접수인", "인계인", "인수인", "특이소견", "비고란")
ABBREVIATION_HEADERS = ("검사위치", "조직검사결과")


def synthetic_record(**changes: object) -> ResultRecord:
    """아래 번호·이름·결과는 테스트를 위한 가짜 값이다."""
    fields = {
        "report_ref": "00001234",
        "chart_no": "00005678",
        "patient_name": "테스트환자A",
        "exam_date": date(2026, 10, 19),
        "age": "45",
        "sex": "F",
        "entrusted_date": date(2026, 10, 20),
        "exam_kind": "SYNTHETIC_EXAM",
        "specimen_count": 2,
        "report_date": date(2026, 10, 22),
        "site_text": "SYNTHETIC_RAW_SITE_DO_NOT_EXPORT",
        "result_text": "SYNTHETIC_RAW_RESULT_DO_NOT_EXPORT",
    }
    fields.update(changes)
    return ResultRecord(**fields)


@pytest.fixture
def policy():
    return load_policy()


@pytest.fixture
def workbook(policy):
    return new_workbook(policy)


def row_values(workbook, row: int = 2) -> dict[str, object]:
    sheet = workbook[SHEET_NAME]
    return {
        header: sheet.cell(row=row, column=column).value
        for column, header in enumerate(HEADERS, start=1)
    }


def set_cell(workbook, row: int, header: str, value: object) -> None:
    workbook[SHEET_NAME].cell(
        row=row, column=HEADERS.index(header) + 1, value=value
    )


def snapshot(workbook) -> tuple[tuple[object, ...], ...]:
    return tuple(workbook[SHEET_NAME].iter_rows(values_only=True))


def test_new_workbook_has_the_exact_seventeen_headers(workbook):
    sheet = workbook[SHEET_NAME]

    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet.max_row == 1
    assert sheet.max_column == 17


def test_new_record_transcribes_fields_and_leaves_manual_work_to_staff(
    workbook, policy
):
    record = synthetic_record()

    result = merge_records(workbook, [record], policy)
    values = row_values(workbook)

    assert result.appended == 1
    assert result.unchanged == 0
    assert workbook[SHEET_NAME].max_row == 2
    assert {header: values[header] for header in MANUAL_HEADERS} == {
        header: None for header in MANUAL_HEADERS
    }
    assert values["검사일"] == record.exam_date
    assert values["접수번호"] == record.report_ref
    assert values["차트번호"] == record.chart_no
    assert values["수진자명"] == record.patient_name
    assert values["나이"] == record.age
    assert values["성별"] == record.sex
    assert values["위탁일"] == record.entrusted_date
    assert values["검사종류"] == record.exam_kind
    assert values["검체수량"] == record.specimen_count
    assert values["결과보고일"] == record.report_date
    assert (record.report_ref, record.chart_no) in result.fingerprints


@pytest.mark.parametrize("include_source_text", [True, False])
def test_abbreviations_always_require_staff_review(
    workbook, policy, include_source_text
):
    record = synthetic_record(
        **({} if include_source_text else {"site_text": None, "result_text": None})
    )

    result = merge_records(workbook, [record], policy)

    for header in ABBREVIATION_HEADERS:
        assert row_values(workbook)[header] == "검토 필요"
        assert any(
            issue.row == 2
            and issue.column == header
            and issue.code == "ABBREVIATION_PENDING"
            for issue in result.issues
        )


def test_identifiers_keep_leading_zeros_after_an_excel_round_trip(
    workbook, policy, tmp_path
):
    record = synthetic_record()
    merge_records(workbook, [record], policy)

    path = save_new_workbook(workbook, tmp_path / "synthetic-ledger.xlsx")
    saved = load_workbook(path)
    try:
        for header, expected in (
            ("접수번호", record.report_ref),
            ("차트번호", record.chart_no),
        ):
            cell = saved[SHEET_NAME].cell(row=2, column=HEADERS.index(header) + 1)
            assert cell.value == expected
            assert cell.data_type == "s"
    finally:
        saved.close()


def test_identical_records_in_one_batch_do_not_create_duplicate_rows(
    workbook, policy
):
    record = synthetic_record()

    result = merge_records(workbook, [record, record], policy)

    assert result.appended == 1
    assert result.unchanged == 1
    assert workbook[SHEET_NAME].max_row == 2
    assert len(result.fingerprints) == 1


def test_rerun_preserves_staff_columns_and_does_not_duplicate_the_row(
    workbook, policy
):
    record = synthetic_record()
    first = merge_records(workbook, [record], policy)
    for header in MANUAL_HEADERS:
        set_cell(workbook, 2, header, f"SYNTHETIC_STAFF_{header}")
    for header in ABBREVIATION_HEADERS:
        set_cell(workbook, 2, header, f"SYNTHETIC_APPROVED_{header}")
    before = snapshot(workbook)

    second = merge_records(
        workbook, [record], policy, fingerprints=first.fingerprints
    )

    assert second.appended == 0
    assert second.unchanged == 1
    assert snapshot(workbook) == before
    assert second.fingerprints == first.fingerprints


@pytest.mark.parametrize(
    "change",
    [
        {"report_date": date(2026, 10, 23)},
        {"patient_name": "테스트환자변경"},
        {"specimen_count": 3},
        {"site_text": "SYNTHETIC_CORRECTED_RAW_SITE"},
        {"result_text": "SYNTHETIC_CORRECTED_RAW_RESULT"},
    ],
)
def test_changed_source_requires_review_and_cannot_overwrite_an_existing_row(
    workbook, policy, change
):
    first = merge_records(workbook, [synthetic_record()], policy)
    set_cell(workbook, 2, "비고란", "SYNTHETIC_STAFF_NOTE")
    before = snapshot(workbook)

    result = merge_records(
        workbook,
        [synthetic_record(**change)],
        policy,
        fingerprints=first.fingerprints,
    )

    assert result.appended == 0
    assert snapshot(workbook) == before
    assert any(issue.row == 2 and issue.code == "SOURCE_CHANGED" for issue in result.issues)
    # 미승인 변경의 fingerprint를 수락하면 다음 실행에서 변경을 놓친다.
    assert result.fingerprints == first.fingerprints


def test_existing_row_without_provenance_is_left_untouched(workbook, policy):
    record = synthetic_record()
    merge_records(workbook, [record], policy)
    set_cell(workbook, 2, "수진자명", "SYNTHETIC_MANUALLY_ENTERED_NAME")
    set_cell(workbook, 2, "비고란", "SYNTHETIC_EXISTING_NOTE")
    before = snapshot(workbook)

    result = merge_records(workbook, [record], policy)

    assert result.appended == 0
    assert snapshot(workbook) == before
    assert any(
        issue.row == 2 and issue.code == "EXISTING_ROW_REVIEW"
        for issue in result.issues
    )


def test_existing_duplicate_keys_fail_without_exposing_patient_values(
    workbook, policy
):
    record = synthetic_record()
    first = merge_records(workbook, [record], policy)
    sheet = workbook[SHEET_NAME]
    sheet.append([cell.value for cell in sheet[2]])
    before = snapshot(workbook)

    with pytest.raises(ValueError) as error:
        merge_records(workbook, [record], policy, fingerprints=first.fingerprints)

    assert snapshot(workbook) == before
    for private_value in (record.report_ref, record.chart_no, record.patient_name):
        assert private_value not in str(error.value)


@pytest.mark.parametrize("header_column", [1, 10, 17])
def test_changed_header_fails_before_mutating_the_workbook(
    workbook, policy, header_column
):
    workbook[SHEET_NAME].cell(row=1, column=header_column, value="SYNTHETIC_CHANGED")
    before = snapshot(workbook)
    record = synthetic_record()

    with pytest.raises(ValueError) as error:
        merge_records(workbook, [record], policy)

    assert snapshot(workbook) == before
    assert record.patient_name not in str(error.value)
    assert record.chart_no not in str(error.value)


@pytest.mark.parametrize(
    "missing_key",
    [{"report_ref": ""}, {"chart_no": ""}, {"report_ref": "  "}, {"chart_no": "  "}],
)
def test_missing_key_fails_without_mutation_or_patient_values_in_errors(
    workbook, policy, missing_key
):
    record = synthetic_record(**missing_key)
    before = snapshot(workbook)

    with pytest.raises(ValueError) as error:
        merge_records(workbook, [record], policy)

    assert snapshot(workbook) == before
    assert record.patient_name not in str(error.value)
    for value in (record.report_ref, record.chart_no):
        if value.strip():
            assert value not in str(error.value)


def test_raw_report_text_is_absent_from_excel_and_diagnostics(
    workbook, policy, caplog, capsys
):
    record = synthetic_record()

    with caplog.at_level("DEBUG"):
        result = merge_records(workbook, [record], policy)
        output = BytesIO()
        workbook.save(output)

    with ZipFile(output) as archive:
        excel_content = b"\n".join(archive.read(name) for name in archive.namelist())
    stdio = capsys.readouterr()
    diagnostics = caplog.text + stdio.out + stdio.err + repr(result)
    for raw_text in (record.site_text, record.result_text):
        assert raw_text.encode("utf-8") not in excel_content
        assert raw_text not in diagnostics


@pytest.mark.parametrize(
    "patient_name", ["=1+1", "+SYNTHETIC", "-SYNTHETIC", "@SYNTHETIC"]
)
def test_formula_like_patient_names_are_literal_text(
    workbook, policy, tmp_path, patient_name
):
    merge_records(workbook, [synthetic_record(patient_name=patient_name)], policy)

    path = save_new_workbook(workbook, tmp_path / "synthetic-formula.xlsx")
    saved = load_workbook(path, data_only=False)
    try:
        cell = saved[SHEET_NAME].cell(row=2, column=HEADERS.index("수진자명") + 1)
        assert cell.value == patient_name
        assert cell.data_type == "s"
    finally:
        saved.close()


def test_save_returns_a_new_readable_workbook(workbook, tmp_path):
    target = tmp_path / "synthetic-new.xlsx"

    result = save_new_workbook(workbook, target)

    assert result == target
    assert target.is_file()
    saved = load_workbook(target)
    try:
        assert tuple(cell.value for cell in saved[SHEET_NAME][1]) == HEADERS
    finally:
        saved.close()


def test_save_never_overwrites_an_existing_file(workbook, tmp_path):
    target = tmp_path / "synthetic-existing.xlsx"
    original = b"SYNTHETIC_EXISTING_FILE_MUST_SURVIVE"
    target.write_bytes(original)

    with pytest.raises(FileExistsError):
        save_new_workbook(workbook, target)

    assert target.read_bytes() == original


@pytest.mark.parametrize("field", ["report_ref", "chart_no", "patient_name", "age", "exam_kind"])
@pytest.mark.parametrize("bad_text", ["SYNTHETIC\x00", "X" * 32768])
def test_invalid_excel_text_rejects_entire_batch_before_any_write(workbook, policy, field, bad_text):
    first = synthetic_record()
    second = replace(synthetic_record(report_ref="TEST-SECOND", chart_no="00009999"), **{field: bad_text})
    before = snapshot(workbook)
    with pytest.raises(ValueError) as error:
        merge_records(workbook, [first, second], policy)
    assert snapshot(workbook) == before
    assert bad_text not in str(error.value)
