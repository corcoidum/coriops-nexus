"""지정 원내 폴더 출력. 지정 폴더 밖·링크·클라우드 경로·덮어쓰기는 거부한다."""

from __future__ import annotations

from datetime import date, timedelta
import os

from openpyxl import load_workbook
import pytest
from fastapi.testclient import TestClient

from app.main import create_app
from app.modules.endoscopy_pathology.manual_dashboard import DashboardError, ManualDashboardService
from ledger_helpers import FixedClock


def make_service(folder, clock=None) -> ManualDashboardService:
    return ManualDashboardService(clock=clock or FixedClock(), export_dir=folder)


def test_export_writes_one_new_file_with_tissue_sheets(tmp_path):
    service = make_service(tmp_path)

    record = service.export()

    files = list(tmp_path.iterdir())
    assert [f.name for f in files] == [record.file_name]
    assert record.file_name == "조직검체장부_20261105_100000.xlsx"
    workbook = load_workbook(files[0])
    # 조직검체 장부 사본이라 CLO(A27)·분변잠혈(A22) 시트는 없다. 각 대장에서 따로 출력한다.
    assert workbook.sheetnames == ["Biopsy"]
    biopsy = workbook["Biopsy"]
    headers = [c.value for c in biopsy[1]]
    assert headers[:17][-3:] == ["조직검사결과", "특이소견", "비고란"]
    assert headers[17:] == ["재검 예정일", "재검 상태"]
    assert biopsy.max_row == 1 + 12
    rows = {row[2]: row for row in biopsy.iter_rows(min_row=2, values_only=True)}
    assert rows["T-0006"][14] == "TA#2(L), H#1"
    assert rows["T-0007"][15] == "HP-/Neu+/Mon++/Atr+/IM+"
    assert rows["T-0009"][18] == "재검 예정"
    assert service.snapshot()["exports"][0].file_name == record.file_name


def test_export_keeps_formula_like_text_as_text(tmp_path):
    service = make_service(tmp_path)
    service.register("BIOPSY", "T-5555", "=HYPERLINK(1)", date(2026, 11, 5), procedure_type="stomach", specimen_count="1")

    path = tmp_path / service.export().file_name

    sheet = load_workbook(path)["Biopsy"]
    cell = next(row[3] for row in sheet.iter_rows(min_row=2) if row[2].value == "T-5555")
    assert cell.value == "=HYPERLINK(1)" and cell.data_type == "s"


def test_export_never_overwrites(tmp_path):
    service = make_service(tmp_path)
    service.export()

    with pytest.raises(DashboardError) as error:
        service.export()  # 같은 시각 → 같은 이름
    assert error.value.status_code == 409
    assert len(list(tmp_path.iterdir())) == 1


@pytest.mark.parametrize("folder", ["", "relative/exports"])
def test_export_requires_a_configured_absolute_folder(folder):
    with pytest.raises(DashboardError) as error:
        make_service(folder).export()
    assert error.value.status_code == 409


def test_export_rejects_missing_link_and_cloud_folders(tmp_path):
    missing = tmp_path / "none"
    real = tmp_path / "real"
    real.mkdir()
    link = tmp_path / "link"
    os.symlink(real, link)
    cloud = tmp_path / "OneDrive - 의원" / "exports"
    cloud.mkdir(parents=True)

    for folder in (missing, link, cloud):
        with pytest.raises(DashboardError) as error:
            make_service(folder).export()
        assert error.value.status_code == 409
    assert list(real.iterdir()) == [] and list(cloud.iterdir()) == []


def test_export_route_requires_csrf_and_redirects(tmp_path):
    service = make_service(tmp_path)
    app = create_app(service)
    with TestClient(app) as client:
        denied = client.post("/ledger/export", data={"csrf_token": "INVALID"})
        done = client.post("/ledger/export", data={"csrf_token": app.state.csrf_token}, follow_redirects=False)
        page = client.get("/ledger/excel?notice=exported").text
    assert denied.status_code == 403
    assert done.status_code == 303
    assert "조직검체장부_20261105_100000.xlsx" in page
    assert len(list(tmp_path.iterdir())) == 1


def test_retention_counts_closed_rows_past_the_policy_without_deleting():
    clock = FixedClock()
    service = make_service("", clock)
    assert service.snapshot()["retention"]["expired"] == 0

    clock.moment += timedelta(days=365 * 6)
    retention = service.snapshot()["retention"]

    assert retention["years"] == 5
    assert retention["expired"] == 2  # 종결 2건(재검 불필요) — 자동 삭제하지 않는다
    assert retention["status"] == "REQUIRES_OWNER_AND_LEGAL_REVIEW"
    assert sum(service.snapshot(ledger=name)["total"] for name in ("tissue", "clo", "fobt")) == 14
