"""분변잠혈 대장(A22): 조직검체 장부에서 Stool을 빼고 따로 보여준다. 합성 자료만 쓴다."""

from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
import re
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.main import create_app
from app.modules.endoscopy_pathology.ledger_db import MemoryStore, replace_entry
from app.modules.endoscopy_pathology.ledger_model import Handoff
from app.modules.endoscopy_pathology.manual_dashboard import DashboardError, ManualDashboardService
from ledger_helpers import FixedClock

DAY = date(2026, 11, 4)
CLOCK = datetime(2026, 11, 5, 1, 0, tzinfo=timezone.utc)  # 2026-11-05(목) 10:00 KST


@pytest.fixture
def service(tmp_path):
    return ManualDashboardService(clock=FixedClock(CLOCK), store=MemoryStore(), export_dir=tmp_path)


@pytest.fixture
def client(tmp_path):
    app = create_app(ManualDashboardService(clock=FixedClock(CLOCK), store=MemoryStore(), export_dir=tmp_path))
    with TestClient(app) as test_client:
        test_client.app_state = app.state
        yield test_client


def post(client, path, **data):
    return client.post(path, data={"csrf_token": client.app_state.csrf_token, **data}, follow_redirects=False)


def stool(service, number: int, day: date = DAY):
    return service.register("STOOL", f"T-{number:04}", f"테스트변{number}", day, "M", "70")


def test_ledgers_split_tissue_and_fobt(service):
    biopsy = service.register("BIOPSY", "T-0001", "테스트조직", DAY, procedure_type="stomach", specimen_count="1")
    fobt = stool(service, 2)
    tissue = service.snapshot()
    fobt_view = service.snapshot(ledger="fobt")
    assert [c.id for c in tissue["cases"]] == [biopsy] and list(tissue["kind_labels"]) == ["BIOPSY"]
    assert [c.id for c in fobt_view["cases"]] == [fobt] and fobt_view["kind_labels"] == {"STOOL": "분변잠혈"}
    assert tissue["pickup"]["pending_cases"] == 1
    assert fobt_view["pickup"]["pending_cases"] == 0 and not fobt_view["has_handoff"]
    assert fobt_view["stages"][fobt] == "RESULT_PENDING"
    with pytest.raises(DashboardError):
        service.snapshot(kind="STOOL")
    with pytest.raises(DashboardError):
        service.snapshot(selected_id=fobt)
    with pytest.raises(DashboardError):
        service.snapshot(ledger="unknown")


def test_fobt_page_and_actions_return_to_fobt(client):
    page = client.get("/dashboard").text
    assert 'href="/fobt"' in page and "분변잠혈 대장" in page
    response = post(client, "/ledger/register", kind="STOOL", exam_date="2026-11-04", chart_no="T-0301",
                    patient_name="테스트변", sex_age="M/70")
    location = response.headers["location"]
    assert response.status_code == 303 and location.startswith("/fobt?")
    fobt_page = client.get(location).text
    assert "분변잠혈 대장" in fobt_page and 'action="/fobt"' in fobt_page
    assert 'class="kind-tabs"' not in fobt_page  # 종류가 하나라 종류 탭이 없다
    assert "T-0301" not in client.get("/dashboard").text
    entry_id = location.split("selected_id=")[1].split("&")[0]
    assert client.get(f"/dashboard?selected_id={entry_id}").status_code == 404

    assert "오늘 수탁기관 인계" not in fobt_page
    receive = post(client, f"/ledger/cases/{entry_id}/receive", result_notation="Negative 6.0")
    assert receive.headers["location"] == f"/fobt?selected_id={entry_id}"
    for href in re.findall(r'href="(/fobt\?[^"]+)"', client.get(receive.headers["location"]).text):
        assert client.get(href.replace("&amp;", "&")).status_code == 200, href


def test_unknown_ledger_in_forms_falls_back_to_tissue(client):
    post(client, "/ledger/register", kind="BIOPSY", exam_date="2026-11-04", chart_no="T-0302",
         patient_name="테스트조직", sex_age="M/70", procedure_type="stomach", specimen_count="1")
    handoff = post(client, "/ledger/handoff", registered_by="테스트접수", handed_by="테스트인계",
                   receiver_name="테스트수령", confirmed_count="1", ledger="//evil.example")
    assert handoff.headers["location"].startswith("/dashboard?")


def test_tissue_copy_export_has_no_stool_sheet(service, tmp_path):
    service.register("BIOPSY", "T-0001", "테스트조직", DAY, procedure_type="stomach", specimen_count="1")
    stool(service, 2)
    record = service.export()
    assert load_workbook(tmp_path / record.file_name).sheetnames == ["Biopsy"]


def test_paper_form_three_lists_only_stool(service):
    service.register("BIOPSY", "T-0001", "테스트조직", DAY, procedure_type="stomach", specimen_count="1")
    entry = stool(service, 2)
    service.receive(entry, "", "Negative 6.0", "")
    data = service.paper_rows("fobt", DAY, DAY)
    assert [row["chart_no"] for row in data["rows"]] == ["T-0002"]
    assert data["rows"][0]["result"] == "Negative 6.0" and data["title"] == "분변잠혈검사 결과기록대장"


def test_paper_form_three_exports_two_blocks_per_page(service, tmp_path):
    for number in range(1, 46):  # 쪽당 40줄(왼쪽 20 + 오른쪽 20) → 2쪽
        stool(service, number)
    record, subjects = service.export_paper("fobt", DAY, DAY)
    sheet = load_workbook(tmp_path / record.file_name).active
    headers = [sheet.cell(3, column).value for column in range(1, 15)]
    assert headers[:7] == headers[7:] and headers[0] == "접수\n일자" and headers[6].startswith("결과\n(참고치")
    assert "[검진] Stool occult blood 정량" in sheet["A2"].value
    assert sheet.cell(4, 2).value == "T-0001" and sheet.cell(23, 2).value == "T-0020"   # 왼쪽 1~20
    assert sheet.cell(4, 9).value == "T-0021" and sheet.cell(23, 9).value == "T-0040"   # 오른쪽 21~40
    assert sheet.cell(24, 2).value == "T-0041" and sheet.cell(28, 2).value == "T-0045"  # 2쪽 왼쪽
    assert sheet.cell(24, 9).value in (None, "")                                      # 남은 칸은 빈 줄
    assert [brk.id for brk in sheet.row_breaks.brk] == [23]
    assert sheet.print_title_rows == "$3:$3" and sheet.page_setup.orientation == "landscape"
    assert len(subjects) == 45


def test_paper_page_offers_form_three(client):
    page = client.get("/ledger/paper?form=fobt&date_from=2026-11-04&date_to=2026-11-04")
    assert page.status_code == 200 and "③ 분변잠혈 대장" in page.text and "접수일자" in page.text


def test_followup_export_is_scoped_by_ledger(service):
    stool(service, 3)
    with pytest.raises(DashboardError) as tissue_only:
        service.export_followups(kind="STOOL", due="due")
    assert tissue_only.value.status_code == 422


def reported(service, number: int, result: str, chart: int | None = None):
    """접수 후 결과를 입력한 분변잠혈 1건(원장 확인 대기)."""
    entry = service.register("STOOL", f"T-{chart or number:04}", f"테스트변{chart or number}", DAY, "M", "70")
    service.receive(entry, "", result, "")
    return entry


def test_fobt_has_no_followup_date_only_doctor_confirmation(service):
    entry = reported(service, 11, "Negative 6.0")
    view = service.snapshot(ledger="fobt")
    assert "FOLLOWUP" not in view["stage_cards"] and view["stage_labels"]["DOCTOR_PENDING"] == "원장 확인 대기"
    with pytest.raises(DashboardError) as no_due:
        service.decide(entry, "required", DAY)
    assert no_due.value.status_code == 422
    with pytest.raises(DashboardError):
        service.snapshot(ledger="fobt", due="due")
    service.decide(entry, "confirmed", reason="", actor="테스트원장")
    view = service.snapshot(ledger="fobt", selected_id=entry)
    assert view["stages"][entry] == "CLOSED" and view["selected"].followup_due is None
    assert view["alerts"][entry].text == "원장 확인"
    # 조직검체는 그대로 재검 날짜를 받고, '확인'만으로는 닫히지 않는다.
    biopsy = service.register("BIOPSY", "T-0012", "테스트조직", DAY, procedure_type="stomach", specimen_count="1")
    assert "FOLLOWUP" in service.snapshot()["stage_cards"]
    with pytest.raises(DashboardError):
        service.decide(biopsy, "confirmed")


def test_positive_results_count_cfs_subjects_by_chart(service):
    first = reported(service, 21, "Positive 152")
    reported(service, 22, "+", chart=21)          # 같은 수진자 → 1명
    reported(service, 23, "양성")
    reported(service, 24, "Negative 6.0")
    reported(service, 25, "152")                  # 양성 표기 없는 수치는 세지 않는다
    view = service.snapshot(ledger="fobt")
    assert view["flags"]["positive_subjects"] == 2 and view["positive_label"] == "CFS 권고 대상"
    assert view["alerts"][first].text == "양성 · CFS 권고 대상"
    only = service.snapshot(ledger="fobt", positive=True)
    assert {c.chart_no for c in only["cases"]} == {"T-0021", "T-0023"} and only["matched"] == 3
    service.decide(first, "confirmed", reason="CFS 안내")
    assert service.snapshot(ledger="fobt")["alerts"][first].text == "CFS 권고 대상"
    with pytest.raises(DashboardError):
        service.snapshot(positive=True)  # 조직검체 장부에는 양성 대상자 목록이 없다


def test_fobt_page_shows_cfs_card_and_confirm_form(client):
    post(client, "/ledger/register", kind="STOOL", exam_date="2026-11-04", chart_no="T-0401",
         patient_name="테스트변", sex_age="M/70")
    entry_id = client.app_state.dashboard_service.snapshot(ledger="fobt", query="T-0401")["cases"][0].id
    post(client, f"/ledger/cases/{entry_id}/receive", result_notation="Positive 152")
    page = client.get(f"/fobt?selected_id={entry_id}").text
    assert "CFS 권고 대상자" in page and 'value="confirmed"' in page and "재검 날짜 지정" not in page
    assert "<span>재검 예정</span>" not in page and "/fobt?stage=ALL&amp;kind=ALL&amp;positive=1" in page
    positive_page = client.get("/fobt?stage=ALL&kind=ALL&positive=1").text
    assert "T-0401" in positive_page and "CFS 시행 여부는 원장이 결정합니다" in positive_page
    bulk = post(client, "/ledger/bulk-decide", entry_ids=str(entry_id), reason="", ledger="fobt")
    assert bulk.status_code == 303 and bulk.headers["location"].startswith("/fobt?")
    done = client.get(f"/fobt?selected_id={entry_id}").text
    assert "원장 확인" in done and "재검 불필요" not in done


def test_paper_form_three_notes_cfs_subjects_in_period(service, tmp_path):
    reported(service, 31, "Positive 152")
    reported(service, 32, "Negative 6.0")
    data = service.paper_rows("fobt", DAY, DAY)
    assert data["summary"] == "CFS 권고 대상자 1명"
    record, _ = service.export_paper("fobt", DAY, DAY)
    assert "CFS 권고 대상자 1명" in load_workbook(tmp_path / record.file_name).active["A2"].value


def test_fobt_intake_has_no_bottle_count_and_is_not_in_tissue_today_card(client):
    page = client.get("/fobt?new=1").text
    assert 'name="specimen_count"' not in page and "개수(병)" not in page
    assert "다음날" in page and "오늘 수탁기관 인계" not in page
    response = post(client, "/ledger/register", kind="STOOL", exam_date="2026-11-05", chart_no="T-0501",
                    patient_name="테스트제출", specimen_count="3")  # 예전 폼의 값도 새 접수에는 저장하지 않는다.
    assert response.status_code == 303
    service = client.app_state.dashboard_service
    view = service.snapshot(ledger="fobt")
    assert view["cases"][0].specimen_count is None and view["counts"]["RESULT_PENDING"] == 1
    home = client.get("/").text
    tissue = re.search(r'aria-labelledby="today-tissue">(.*?)</section>', home, re.S).group(1)
    assert "분변잠혈" not in tissue and "<strong>0</strong>" in tissue
    assert "당일 제출 대상 접수 → 다음날 결과 입력 → 원장 확인" in home


def test_fobt_submitted_today_can_receive_tomorrows_result_without_handoff(service):
    entry_id = stool(service, 502, date(2026, 11, 5))
    assert service.snapshot(ledger="fobt")["stages"][entry_id] == "RESULT_PENDING"
    service._clock.moment += timedelta(days=1)
    service.receive(entry_id, "", "Negative 6.0", "")
    view = service.snapshot(ledger="fobt", selected_id=entry_id)
    assert view["selected"].reported_on == date(2026, 11, 6)
    assert view["selected"].handoff_id is None and view["stages"][entry_id] == "DOCTOR_PENDING"
    service.decide(entry_id, "confirmed", actor="테스트원장")
    assert service.snapshot(ledger="fobt")["stages"][entry_id] == "CLOSED"


def test_old_fobt_handoff_form_cannot_handoff_tissue_cases(client):
    service = client.app_state.dashboard_service
    biopsy = service.register("BIOPSY", "T-0505", "테스트조직인계", DAY,
                              procedure_type="stomach", specimen_count="1")
    fobt = stool(service, 506)
    response = post(client, "/ledger/handoff", registered_by="테스트접수", handed_by="테스트인계",
                    receiver_name="테스트수령", confirmed_count="1", ledger="fobt")
    assert response.status_code == 422
    assert service.snapshot()["stages"][biopsy] == "HANDOFF_PENDING"
    assert service.snapshot(ledger="fobt")["stages"][fobt] == "RESULT_PENDING"
    assert service._store.handoffs() == []


def test_legacy_mixed_handoff_is_preserved_but_tissue_count_excludes_stool(service):
    biopsy = service.register("BIOPSY", "T-0503", "테스트조직", DAY, procedure_type="stomach", specimen_count="3")
    fobt = stool(service, 504)
    entries = {e.id: e for e in service._store.entries()}
    handoff = Handoff(uuid4(), service._clock.now(), "테스트접수", "테스트인계", "테스트수령", 4, 2)
    service._store.write(handoff=handoff, updated=[
        replace_entry(entries[biopsy], service._clock.now(), handoff_id=handoff.id, entrusted_on=DAY),
        replace_entry(entries[fobt], service._clock.now(), handoff_id=handoff.id, entrusted_on=DAY, specimen_count=1),
    ])
    [shown] = service.snapshot()["pickup"]["handed_today"]
    assert (shown.case_count, shown.specimen_count) == (1, 3)
    [stored] = service._store.handoffs()
    assert (stored.case_count, stored.specimen_count) == (2, 4)
    assert next(e for e in service._store.entries() if e.id == fobt).specimen_count == 1
