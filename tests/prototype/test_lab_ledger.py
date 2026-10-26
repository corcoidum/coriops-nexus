"""합성 데이터 전용 혈액검체 장부 계약: 실데이터·외부 조회를 사용하지 않는다."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.modules.screening_lab.ledger import (
    LabItem,
    LabRecord,
    format_result,
    load_policy,
    merge_records,
    new_workbook,
    save_new_workbook,
)
from app.modules.screening_lab.ledger_demo import synthetic_records


SHEET_NAME = "혈액검체장부"
HEADERS = (
    "날짜",
    "접수번호",
    "차트번호",
    "성명",
    "성별/나이",
    "분류(일반/간)",
    "접수인(서명)",
    "위탁날짜",
    "인계인(서명)",
    "수령인(서명)",
    "회신날짜",
    "결과",
)
TITLE = "혈액병리검체 접수, 인계, 인수 및 회신결과 대장 (일반<6종.9종>/간)"
FIRST_ROW = 3  # 1행 제목, 2행 열 이름
MANUAL_HEADERS = ("접수인(서명)", "인계인(서명)", "수령인(서명)")


def synthetic_record(**changes: object) -> LabRecord:
    """아래 번호·이름·수치는 테스트를 위한 가짜 값이다."""
    fields = {
        "report_ref": "00004321",
        "chart_no": "00008765",
        "patient_name": "테스트수진자A",
        "received_date": date(2026, 10, 19),
        "sex": "F",
        "age": "52",
        "exam_kinds": ("일반검진9종", "간검진(AFP)"),
        "entrusted_date": date(2026, 10, 19),
        "report_date": date(2026, 10, 21),
        "items": (
            LabItem("AST", "52", "H"),
            LabItem("ALT", "38", None),
            LabItem("HDL", "31", "L"),
        ),
    }
    fields.update(changes)
    return LabRecord(**fields)


@pytest.fixture
def policy():
    return load_policy()


@pytest.fixture
def workbook(policy):
    return new_workbook(policy)


def row_values(workbook, row: int = FIRST_ROW) -> dict[str, object]:
    sheet = workbook[SHEET_NAME]
    return {header: sheet.cell(row=row, column=column).value for column, header in enumerate(HEADERS, start=1)}


def test_headers_follow_user_order(workbook, policy):
    assert policy.headers == HEADERS
    sheet = workbook[SHEET_NAME]
    assert tuple(cell.value for cell in sheet[2]) == HEADERS
    assert sheet["A1"].value == TITLE and "A1:L1" in {str(r) for r in sheet.merged_cells.ranges}


def test_one_row_holds_whole_patient_result(workbook, policy):
    result = merge_records(workbook, [synthetic_record()], policy)
    values = row_values(workbook)
    assert result.appended == 1
    assert workbook[SHEET_NAME].max_row == FIRST_ROW
    assert values["분류(일반/간)"] == "일반검진9종, 간검진(AFP)"
    assert values["접수번호"] == "00004321"
    assert values["날짜"] == date(2026, 10, 19)


def test_result_lists_only_flagged_items(workbook, policy):
    merge_records(workbook, [synthetic_record()], policy)
    assert row_values(workbook)["결과"] == "AST 52 H, HDL 31 L"


def test_all_normal_writes_normal_label(policy):
    assert format_result([LabItem("ALT", "38", None)], policy.normal_label) == "정상"


def test_no_items_leaves_result_blank(policy):
    assert format_result([], policy.normal_label) is None


def test_all_normal_row_shows_normal(workbook, policy):
    merge_records(workbook, [synthetic_record(items=(LabItem("ALT", "38"),))], policy)
    assert row_values(workbook)["결과"] == "정상"


def test_unregistered_item_name_is_written_and_reviewed(workbook, policy):
    record = synthetic_record(items=(LabItem("SGOT", "52", "H"), LabItem("미등록정상", "1")))
    result = merge_records(workbook, [record], policy)
    assert row_values(workbook)["결과"] == "SGOT 52 H"
    assert [issue.code for issue in result.issues] == ["ITEM_NAME_UNREGISTERED"]


def test_registered_item_names_follow_user_list(policy):
    assert policy.item_names == ("Hb", "Glu", "CHOL", "TG", "HDL", "LDL", "AST", "ALT", "GGT", "Cr")


def test_manual_columns_stay_blank(workbook, policy):
    merge_records(workbook, [synthetic_record()], policy)
    values = row_values(workbook)
    assert all(values[header] is None for header in MANUAL_HEADERS)


def test_staff_edits_are_preserved_on_rerun(workbook, policy):
    first = merge_records(workbook, [synthetic_record()], policy)
    workbook[SHEET_NAME].cell(FIRST_ROW, HEADERS.index("수령인(서명)") + 1, "직원A")
    again = merge_records(workbook, [synthetic_record()], policy, fingerprints=first.fingerprints)
    assert again.appended == 0 and again.unchanged == 1
    assert row_values(workbook)["수령인(서명)"] == "직원A"


def test_critical_value_creates_report_issue_and_highlight(workbook, policy):
    record = synthetic_record(items=(LabItem("Glu", "999", "H", critical=True),))
    result = merge_records(workbook, [record], policy)
    column = HEADERS.index("결과") + 1
    assert [issue.code for issue in result.issues] == ["CRITICAL_REPORT"]
    assert result.issues[0].row == FIRST_ROW
    assert workbook[SHEET_NAME].cell(FIRST_ROW, column).fill.fgColor.rgb.endswith("F8D7DA")


def test_non_critical_row_has_no_issue(workbook, policy):
    result = merge_records(workbook, [synthetic_record()], policy)
    assert result.issues == ()


def test_changed_source_is_reviewed_not_overwritten(workbook, policy):
    first = merge_records(workbook, [synthetic_record()], policy)
    changed = synthetic_record(items=(LabItem("AST", "70", "H"),))
    again = merge_records(workbook, [changed], policy, fingerprints=first.fingerprints)
    assert [issue.code for issue in again.issues] == ["SOURCE_CHANGED"]
    assert row_values(workbook)["결과"] == "AST 52 H, HDL 31 L"


def test_existing_row_without_fingerprint_is_reviewed(workbook, policy):
    merge_records(workbook, [synthetic_record()], policy)
    again = merge_records(workbook, [synthetic_record()], policy)
    assert [issue.code for issue in again.issues] == ["EXISTING_ROW_REVIEW"]


@pytest.mark.parametrize(
    ("changes", "code"),
    [
        ({"exam_kinds": ("종합검진",)}, "UNKNOWN_EXAM_KIND"),
        ({"exam_kinds": ("일반검진6종", "일반검진6종")}, "DUPLICATE_EXAM_KIND"),
        ({"items": (LabItem("AST", "52", "HH"),)}, "UNKNOWN_RESULT_FLAG"),
        ({"items": (LabItem("", "52", "H"),)}, "INVALID_ITEM_TEXT"),
        ({"chart_no": " "}, "MISSING_RECORD_KEY_OR_NAME"),
        ({"report_date": "2026-10-21"}, "INVALID_DATE_TYPE"),
        ({"exam_kinds": ["일반검진6종"]}, "INVALID_COLLECTION_TYPE"),
    ],
)
def test_invalid_input_rejects_whole_batch(workbook, policy, changes, code):
    with pytest.raises(ValueError, match=code):
        merge_records(workbook, [synthetic_record(report_ref="OK-1"), synthetic_record(**changes)], policy)
    assert workbook[SHEET_NAME].max_row == 2


def test_duplicate_key_in_batch_rejected(workbook, policy):
    with pytest.raises(ValueError, match="DUPLICATE_BATCH_KEY"):
        merge_records(workbook, [synthetic_record(), synthetic_record()], policy)


def test_formula_like_text_is_stored_as_text(workbook, policy):
    merge_records(workbook, [synthetic_record(patient_name="=테스트()")], policy)
    content = BytesIO()
    workbook.save(content)
    reopened = load_workbook(BytesIO(content.getvalue()))
    cell = reopened[SHEET_NAME].cell(FIRST_ROW, HEADERS.index("성명") + 1)
    assert cell.data_type == "s" and cell.value == "=테스트()"


def test_header_mismatch_rejected(workbook, policy):
    workbook[SHEET_NAME].cell(2, 1, "검사일")
    with pytest.raises(ValueError, match="HEADER_MISMATCH"):
        merge_records(workbook, [synthetic_record()], policy)


def test_save_never_overwrites(tmp_path, workbook):
    target = tmp_path / "lab.xlsx"
    save_new_workbook(workbook, target)
    before = target.read_bytes()
    with pytest.raises(FileExistsError):
        save_new_workbook(workbook, target)
    assert target.read_bytes() == before


def test_record_repr_hides_patient_values():
    assert "테스트수진자A" not in repr(synthetic_record())
    # 기본 repr의 메모리 주소(16진수)와 겹치지 않도록 점이 들어간 값을 쓴다.
    assert "52.7" not in repr(LabItem("AST", "52.7", "H"))


def test_demo_records_are_synthetic_and_idempotent(workbook, policy):
    records = synthetic_records()
    assert all(record.patient_name.startswith("테스트") for record in records)
    first = merge_records(workbook, records, policy)
    again = merge_records(workbook, records, policy, fingerprints=first.fingerprints)
    assert first.appended == 10 and again.appended == 0
    assert not any(issue.code == "ITEM_NAME_UNREGISTERED" for issue in first.issues)
    assert sum(issue.code == "CRITICAL_REPORT" for issue in first.issues) == 2


def test_lab_module_does_not_import_pathology_module():
    source = Path(__file__).resolve().parents[2] / "app/modules/screening_lab"
    for path in source.glob("*.py"):
        assert "endoscopy_pathology" not in path.read_text(encoding="utf-8")


def test_policy_rejects_wrong_column_count(tmp_path, policy):
    broken = tmp_path / "broken.yaml"
    broken.write_text(
        "sheet_name: x\nfreeze_panes: A2\nexam_kinds: []\nresult_flags: [H]\n"
        "columns:\n  - {header: 접수일, field: received_date, kind: date, width: 10}\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="INVALID_LEDGER_POLICY"):
        load_policy(broken)


@pytest.mark.parametrize(("sex", "age", "expected"), [("F", "52", "F/52"), ("M", None, "M/"), (None, None, None)])
def test_sex_and_age_share_one_cell(workbook, policy, sex, age, expected):
    merge_records(workbook, [synthetic_record(sex=sex, age=age)], policy)
    assert row_values(workbook)["성별/나이"] == expected


def test_title_row_survives_reopen_and_rows_start_below_headers(workbook, policy):
    merge_records(workbook, [synthetic_record(), synthetic_record(report_ref="00004322")], policy)
    content = BytesIO()
    workbook.save(content)
    sheet = load_workbook(BytesIO(content.getvalue()))[SHEET_NAME]
    assert sheet["A1"].value == TITLE
    assert [sheet.cell(row, 2).value for row in (3, 4)] == ["00004321", "00004322"]
    assert sheet.freeze_panes == "E3"
    # 종이 출력용: 가로 A4 한 쪽 폭, 쪽마다 제목·열 이름 반복
    assert sheet.page_setup.orientation == "landscape" and sheet.page_setup.fitToWidth == 1
    assert sheet.print_title_rows == "$1:$2"
