"""종이 대장(A15)과 맞춘 접수·대조·출력. ① 검체 관리 장부, ② 조직병리검체 접수·인계·인수·회신결과 대장."""

from __future__ import annotations

from datetime import date, datetime, timezone
import re

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.main import create_app
from app.modules.endoscopy_pathology.excel_io import PAPER_FORMS
from app.modules.endoscopy_pathology.ledger_db import MemoryStore
from app.modules.endoscopy_pathology.manual_dashboard import (
    DashboardError, ManualDashboardService, parse_sex_age,
)
from ledger_helpers import FixedClock

# 2026-11-05(목) 10:00 KST
TODAY = date(2026, 11, 5)
DAY = date(2026, 11, 4)


@pytest.fixture
def service(tmp_path):
    return ManualDashboardService(clock=FixedClock(datetime(2026, 11, 5, 1, 0, tzinfo=timezone.utc)),
                                  store=MemoryStore(), export_dir=tmp_path)


def endoscopy(service, chart="T-0101", name="테스트위내시경", procedure="stomach", biopsy=True, rut=True,
              count="2", site="antrum", sex_age="M/65", day=DAY):
    return service.register_endoscopy(chart, name, day, sex_age, procedure, biopsy, site, count, rut)


@pytest.mark.parametrize(("text", "expected"), [
    ("M/65", ("M", "65")), ("F/56", ("F", "56")), (" m / 7 ", ("M", "7")), ("남/70", ("M", "70")),
    ("여/48", ("F", "48")), ("", ("", "")), ("M", ("M", "")), ("/65", ("", "65")),
])
def test_sex_age_is_read_as_written_on_the_paper_ledger(text, expected):
    assert parse_sex_age(text) == expected


@pytest.mark.parametrize("text", ["X/65", "M/abc", "65세", "M/1234", "M/65/1"])
def test_unclear_sex_age_is_refused_not_guessed(text):
    with pytest.raises(DashboardError):
        parse_sex_age(text)


def test_one_paper_row_creates_biopsy_and_rut_together(service):
    ids = endoscopy(service)
    entries = {e.id: e for e in service._store.entries()}
    biopsy, rut = entries[ids[0]], entries[ids[1]]
    assert (biopsy.kind, biopsy.specimen_count, biopsy.site, biopsy.sex, biopsy.age) == ("BIOPSY", 2, "antrum", "M", 65)
    assert (rut.kind, rut.procedure_type, rut.specimen_count, rut.sex, rut.age) == ("CLO", "stomach", None, "M", 65)
    stages = service.snapshot()["stages"]
    # 조직검사는 오늘 인계 대상, RUT는 원내 판정이라 바로 결과 대기
    assert (stages[biopsy.id], stages[rut.id]) == ("HANDOFF_PENDING", "RESULT_PENDING")


@pytest.mark.parametrize(("changes", "message"), [
    ({"biopsy": False, "rut": False}, "하나 이상"),
    ({"procedure": "colon", "rut": True}, "stomach"),
    ({"procedure": ""}, "검사명"),
    ({"count": ""}, "개수(병)"),
    ({"sex_age": "남자65"}, "M/65"),
])
def test_invalid_paper_row_saves_nothing(service, changes, message):
    with pytest.raises(DashboardError) as error:
        endoscopy(service, **changes)
    assert message in str(error.value)
    assert service._store.entries() == []


def test_specimen_ledger_groups_one_endoscopy_per_row(service):
    endoscopy(service)
    endoscopy(service, chart="T-0102", name="테스트대장", procedure="colon", rut=False, count="3", site="sigmoid")
    endoscopy(service, chart="T-0103", name="테스트RUT만", biopsy=False, sex_age="F/56")
    endoscopy(service, chart="T-0104", name="테스트다른날", day=date(2026, 11, 3))

    data = service.paper_rows("specimen", DAY, DAY)
    rows = {r["chart_no"]: r for r in data["rows"]}
    assert list(rows) == ["T-0101", "T-0102", "T-0103"]
    assert (rows["T-0101"]["biopsy_text"], rows["T-0101"]["rut_done"], rows["T-0101"]["rut_result"]) == \
        ("Y · antrum · 2병", "Y", "판정 전")
    assert (rows["T-0102"]["procedure"], rows["T-0102"]["biopsy_text"], rows["T-0102"]["rut_done"]) == \
        ("colon", "Y · sigmoid · 3병", "N")
    assert (rows["T-0103"]["biopsy_text"], rows["T-0103"]["rut_done"]) == ("N", "Y")

    rut_id = next(e.id for e in service._store.entries() if e.kind == "CLO" and e.chart_no == "T-0101")
    service.receive(rut_id, "", "+", "")
    assert service.paper_rows("specimen", DAY, DAY)["rows"][0]["rut_result"] == "+"


def test_pathology_ledger_follows_handoff_and_result(service):
    endoscopy(service)
    endoscopy(service, chart="T-0102", name="테스트대장", procedure="colon", rut=False, count="1", sex_age="F/56")
    before = service.paper_rows("pathology", DAY, DAY)["rows"]
    assert [r["chart_no"] for r in before] == ["T-0101", "T-0102"]  # RUT는 ② 대장에 없다
    assert [r["sex_age"] for r in before] == ["M/65", "F/56"]
    assert before[0]["entrusted_on"] is None and before[0]["received_by"] == ""

    service.handoff("테스트 접수", "테스트 인계", "테스트 수령", 3)
    biopsy_id = next(e.id for e in service._store.entries() if e.kind == "BIOPSY" and e.chart_no == "T-0101")
    service.receive(biopsy_id, "T9900101", "TA#2(L)", "antrum", "HP-/Atr+")
    row = service.paper_rows("pathology", DAY, DAY)["rows"][0]
    assert (row["registered_by"], row["handed_by"], row["received_by"], row["entrusted_on"]) == \
        ("테스트 접수", "테스트 인계", "테스트 수령", TODAY)
    assert (row["report_ref"], row["reported_on"]) == ("T9900101", TODAY)
    assert row["result"].startswith("TA#2(L)") and "HP-" in row["result"]


@pytest.mark.parametrize(("form", "start", "end"), [
    ("unknown", None, None), ("specimen", date(2026, 11, 5), date(2026, 11, 1)),
    ("specimen", date(2023, 1, 1), date(2026, 11, 5)),
])
def test_paper_view_rejects_bad_requests(service, form, start, end):
    with pytest.raises(DashboardError):
        service.paper_rows(form, start, end)


@pytest.mark.parametrize("form", ["specimen", "pathology"])
def test_export_uses_the_paper_layout(service, tmp_path, form):
    endoscopy(service)
    record, subjects = service.export_paper(form, DAY, DAY)
    assert subjects == ["T-0101"]
    sheet = load_workbook(tmp_path / record.file_name).active
    spec = PAPER_FORMS[form]
    assert sheet["A1"].value == spec["title"]
    assert [sheet.cell(3, i).value for i in range(1, len(spec["headers"]) + 1)] == list(spec["headers"])
    assert sheet.print_title_rows == "$3:$3" and sheet.page_setup.orientation == "landscape"
    values = [c.value for c in sheet[4]]
    assert "테스트위내시경" in values  # 출력 파일은 종이 대장처럼 전체 성명
    assert sheet.max_row == 4
    assert record.file_name.startswith(spec["file_prefix"] + "_20261104-20261104_")
    assert service.snapshot()["exports"][-1].file_name == record.file_name


def test_export_with_no_rows_is_refused(service):
    with pytest.raises(DashboardError) as error:
        service.export_paper("specimen", DAY, DAY)
    assert error.value.status_code == 422


# ---------- 화면: 실제 링크를 따라가 본다


@pytest.fixture
def client(tmp_path):
    app = create_app(ManualDashboardService(
        clock=FixedClock(datetime(2026, 11, 5, 1, 0, tzinfo=timezone.utc)), store=MemoryStore(), export_dir=tmp_path))
    with TestClient(app) as test_client:
        test_client.app_state = app.state
        yield test_client


def post(client, path, **data):
    return client.post(path, data={"csrf_token": client.app_state.csrf_token, **data}, follow_redirects=False)


def test_intake_form_matches_paper_ledger_and_links_work(client):
    page = client.get("/dashboard?new=1").text
    for label in ("시행날짜", "차트번호", "성명", "성별/나이", "검사명", "조직검사 유무, 위치 및 개수", "부위", "개수(병)"):
        assert label in page
    assert 'name="rut"' not in page and 'href="/clo?new=1"' in page
    # 분변잠혈 접수는 분변잠혈 대장에서만 한다(A22).
    assert 'value="STOOL"' not in page
    stool = client.get("/fobt?new=1")
    assert stool.status_code == 200 and 'value="STOOL"' in stool.text and "접수일자" in stool.text

    response = post(client, "/ledger/register", kind="ENDOSCOPY", exam_date="2026-11-04", chart_no="T-0201",
                    patient_name="테스트화면", sex_age="F/56", procedure_type="stomach", biopsy="on", site="body",
                    specimen_count="1", rut="on")
    assert response.status_code == 303 and "notice=registered" in response.headers["location"]
    assert client.get(response.headers["location"]).status_code == 200

    stool_post = post(client, "/ledger/register", kind="STOOL", exam_date="2026-11-04", chart_no="T-0202",
                      patient_name="테스트변", sex_age="M/70")
    assert stool_post.status_code == 303
    service = client.app_state.dashboard_service
    kinds = sorted((e.kind, e.sex, e.age) for e in service._store.entries())
    # 예전 폼의 rut 값이 함께 와도 CLO를 중복 접수하지 않는다. CLO는 전용 화면에서 접수한다.
    assert kinds == [("BIOPSY", "F", 56), ("STOOL", "M", 70)]
    clo_post = post(client, "/ledger/register", kind="CLO", exam_date="2026-11-04", chart_no="T-0201",
                    patient_name="테스트화면", sex_age="F/56")
    assert clo_post.status_code == 303 and clo_post.headers["location"].startswith("/clo?")
    assert sum(e.kind == "CLO" for e in service._store.entries()) == 1

    bad = post(client, "/ledger/register", kind="ENDOSCOPY", exam_date="2026-11-04", chart_no="T-0203",
               patient_name="테스트없음", procedure_type="colon")
    assert bad.status_code == 422


def test_paper_page_tabs_period_and_export(client):
    post(client, "/ledger/register", kind="ENDOSCOPY", exam_date="2026-11-05", chart_no="T-0301",
         patient_name="테스트대조", sex_age="M/61", procedure_type="stomach", biopsy="on", site="antrum",
         specimen_count="2", rut="on")
    page = client.get("/ledger/paper")
    assert page.status_code == 200
    assert "검체 관리 장부" in page.text and "테***조" in page.text
    assert "테스트대조" not in page.text  # 화면은 성명을 가린다
    links = set(re.findall(r'href="(/ledger/paper[^"]*)"', page.text)) | set(
        re.findall(r'href="(/[^"]*)"', client.get("/dashboard").text))
    for link in links:
        assert client.get(link.replace("&amp;", "&")).status_code == 200, link

    pathology = client.get("/ledger/paper?form=pathology&date_from=2026-11-05&date_to=2026-11-05").text
    for header in PAPER_FORMS["pathology"]["headers"]:
        assert header in pathology
    assert "M/61" in pathology

    exported = post(client, "/ledger/paper/export", form="pathology", date_from="2026-11-05", date_to="2026-11-05")
    assert exported.status_code == 303 and "notice=paper_exported" in exported.headers["location"]
    assert "종이 대장 양식으로" in client.get(exported.headers["location"]).text
    empty = post(client, "/ledger/paper/export", form="specimen", date_from="2026-10-01", date_to="2026-10-01")
    assert empty.status_code == 422
