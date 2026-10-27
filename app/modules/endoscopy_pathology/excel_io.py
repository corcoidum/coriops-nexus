"""기존 엑셀 장부(Biopsy·CLO·Stool) 읽기와 지정 폴더 출력.

- 업로드 파일은 메모리에서만 읽고 디스크에 저장하지 않는다.
- 오류 보고에는 시트·행 번호·열 이름만 담고 셀 값을 담지 않는다.
- 비고 등 무시 열은 저장·표시하지 않는다(값 대신 헤더 이름만 보고).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import yaml

from .ledger_model import KINDS, KST, LedgerEntry
from .result_notation import NotationError, ResultCodeBook


DEFAULT_CONFIG = Path(__file__).resolve().parents[3] / "config" / "ledger_app.yaml"
CHART_NO = re.compile(r"^[0-9A-Za-z-]{1,20}$")
REPORT_REF = re.compile(r"^[0-9A-Za-z-]{4,20}$")
DATE_FORMATS = ("%Y.%m.%d", "%Y-%m-%d", "%Y/%m/%d", "%Y%m%d")
FIELD_LABELS = {
    "exam_date": "검사일", "report_ref": "접수번호", "chart_no": "차트번호", "patient_name": "수진자명",
    "age": "나이", "sex": "성별", "entrusted_on": "위탁일", "registered_by": "접수인",
    "procedure_type": "검사종류", "site": "검사위치", "specimen_count": "검체수량", "handed_by": "인계인",
    "received_by": "인수인", "reported_on": "결과보고일", "result": "결과", "sydney": "특이소견",
}
ISSUE_LABELS = {
    "REQUIRED_MISSING": "필수 값 없음(가져오지 않음)",
    "INVALID_VALUE": "형식 오류(빈 값으로 가져옴)",
    "INVALID_REQUIRED": "필수 값 형식 오류(가져오지 않음)",
    "NAME_MISMATCH": "같은 차트번호에 다른 이름(가져오지 않음)",
    "MISSING_HEADER": "필수 열 없음(시트 전체 건너뜀)",
    "TOO_MANY_ROWS": "행 수 초과(시트 전체 건너뜀)",
    "DUPLICATE_IN_FILE": "파일 안 중복 행(첫 행만 사용)",
}


class WorkbookError(ValueError):
    """사용자에게 보여줄 수 있는 파일 단위 오류. 셀 값을 담지 않는다."""


@dataclass(frozen=True)
class LedgerAppConfig:
    kinds: dict[str, dict]
    fields: dict[str, tuple[str, ...]]
    required_fields: tuple[str, ...]
    ignored_headers: frozenset[str]
    procedure_types: dict[str, str]
    sexes: dict[str, str]
    short_result_pattern: re.Pattern
    max_rows_per_sheet: int
    export: dict

    @classmethod
    def load(cls, path: Path | None = None) -> LedgerAppConfig:
        with open(path or DEFAULT_CONFIG, encoding="utf-8") as source:
            raw = yaml.safe_load(source)
        if set(raw["kinds"]) != set(KINDS):
            raise ValueError("INVALID_LEDGER_APP_CONFIG")
        config = cls(
            kinds=raw["kinds"], fields={k: tuple(v) for k, v in raw["fields"].items()},
            required_fields=tuple(raw["required_fields"]), ignored_headers=frozenset(raw["ignored_headers"]),
            procedure_types=raw["procedure_types"], sexes=raw["sexes"],
            short_result_pattern=re.compile(raw["short_result_pattern"]),
            max_rows_per_sheet=int(raw["max_rows_per_sheet"]), export=raw["export"],
        )
        config.check_export_headers()
        return config

    @property
    def handoff_kinds(self) -> frozenset[str]:
        return frozenset(kind for kind, meta in self.kinds.items() if meta["handoff"])

    def check_export_headers(self) -> None:
        for kind, meta in self.kinds.items():
            unknown = set(meta.get("export_headers") or ()) - set(EXPORT_COLUMNS)
            if unknown:
                raise ValueError(f"INVALID_EXPORT_HEADERS:{kind}")

    @property
    def review_kinds(self) -> frozenset[str]:
        """결과 뒤 원장 지시 단계가 있는 종류(A18: CLO는 없음)."""
        return frozenset(kind for kind, meta in self.kinds.items() if meta["doctor_review"])

    def ledger_labels(self) -> dict[str, str]:
        return {kind: self.kinds[kind]["ledger"] for kind in KINDS}

    def header_to_field(self) -> dict[str, str]:
        return {alias: name for name, aliases in self.fields.items() for alias in aliases}


@dataclass(frozen=True)
class Issue:
    sheet: str
    row: int  # 0 = 시트 전체
    code: str
    column: str = ""

    @property
    def label(self) -> str:
        return ISSUE_LABELS.get(self.code, self.code)


@dataclass
class RowDraft:
    """엑셀 1행을 정규화한 값. 저장 전 미리보기·대조에만 쓴다."""

    kind: str
    sheet: str
    row: int
    chart_no: str
    patient_name: str
    exam_date: date
    sex: str | None = None
    age: int | None = None
    procedure_type: str | None = None
    specimen_count: int | None = None
    site: str = ""
    report_ref: str = ""
    registered_by: str = ""
    handed_by: str = ""
    received_by: str = ""
    entrusted_on: date | None = None
    reported_on: date | None = None
    findings: tuple = ()
    sydney: dict = field(default_factory=dict)
    result_code: str = ""
    review_codes: list[str] = field(default_factory=list)

    @property
    def result_notation(self) -> str:
        return ", ".join(f.display() for f in self.findings) if self.findings else self.result_code

    def normalized(self) -> tuple:
        """가져오기 digest용. 행 위치·시트 이름은 넣지 않는다."""
        return (self.kind, self.chart_no, self.exam_date.isoformat(), self.report_ref, self.procedure_type,
                self.specimen_count, self.reported_on and self.reported_on.isoformat(), self.result_notation,
                tuple(sorted(self.sydney.items())))


@dataclass
class SheetResult:
    kind: str
    sheet: str
    rows: list[RowDraft] = field(default_factory=list)
    ignored_headers: list[str] = field(default_factory=list)
    unknown_headers: list[str] = field(default_factory=list)
    skipped: bool = False


@dataclass
class ParsedWorkbook:
    sheets: list[SheetResult]
    issues: list[Issue]
    other_sheets: list[str]


# ---------- 값 변환

def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return ILLEGAL_CHARACTERS_RE.sub("", str(value)).strip()


def _date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value).rstrip(".")
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError


def _int(value, low: int, high: int) -> int | None:
    text = _text(value)
    if not text:
        return None
    number = int(text)
    if not low <= number <= high:
        raise ValueError
    return number


class LedgerReader:
    def __init__(self, config: LedgerAppConfig, codebook: ResultCodeBook):
        self.config = config
        self.codebook = codebook

    def read(self, data: bytes, today: date, kind: str | None = None) -> ParsedWorkbook:
        """kind를 주면 시트 이름과 관계없이 모든 시트를 그 대장으로 읽는다(대장별 파일, A18).
        필수 열이 없는 시트(요약표 등)는 건너뛴 시트로만 보고한다."""
        try:
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, OSError, ValueError, TypeError):
            raise WorkbookError("엑셀(.xlsx) 파일을 읽을 수 없습니다. 암호가 걸려 있으면 해제한 사본을 올려 주세요.") from None
        sheet_kinds = {name.casefold(): kind for kind, meta in self.config.kinds.items() for name in meta["sheet_names"]}
        result = ParsedWorkbook(sheets=[], issues=[], other_sheets=[])
        try:
            for worksheet in workbook.worksheets:
                sheet_kind = kind or sheet_kinds.get(worksheet.title.strip().casefold())
                if sheet_kind is None:
                    result.other_sheets.append(worksheet.title)
                    continue
                sheet = self._read_sheet(worksheet, sheet_kind, today, result.issues, optional=kind is not None)
                if sheet is None:
                    result.other_sheets.append(worksheet.title)
                    continue
                result.sheets.append(sheet)
        finally:
            workbook.close()
        if not result.sheets:
            if kind:
                raise WorkbookError("필수 열(검사일·차트번호·수진자명)이 있는 시트를 찾지 못했습니다. 첫 줄의 열 이름을 확인해 주세요.")
            raise WorkbookError("Biopsy·CLO·Stool 시트를 찾지 못했습니다. 시트 이름을 확인하거나 대장 종류를 골라 주세요.")
        self._check_names_in_file(result)
        return result

    def _read_sheet(self, worksheet, kind: str, today: date, issues: list[Issue],
                    optional: bool = False) -> SheetResult | None:
        sheet = SheetResult(kind=kind, sheet=worksheet.title)
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, ())
        mapping = self.config.header_to_field()
        columns: dict[int, str] = {}
        for index, raw in enumerate(header):
            name = _text(raw)
            if not name:
                continue
            if name in self.config.ignored_headers:
                sheet.ignored_headers.append(name)
            elif name in mapping:
                columns.setdefault(index, mapping[name])
            else:
                sheet.unknown_headers.append(name)
        present = set(columns.values())
        missing = [FIELD_LABELS[f] for f in self.config.required_fields if f not in present]
        if missing:
            if optional:
                return None
            issues.extend(Issue(sheet.sheet, 0, "MISSING_HEADER", label) for label in missing)
            sheet.skipped = True
            return sheet
        seen: set[tuple] = set()
        for row_no, values in enumerate(rows, start=2):
            if row_no - 1 > self.config.max_rows_per_sheet:
                issues.append(Issue(sheet.sheet, 0, "TOO_MANY_ROWS"))
                sheet.rows.clear()
                sheet.skipped = True
                return sheet
            cells = {name: values[index] for index, name in columns.items() if index < len(values)}
            if all(_text(value) == "" for value in cells.values()):
                continue
            draft = self._row(kind, sheet.sheet, row_no, cells, today, issues)
            if draft is None:
                continue
            key = (draft.report_ref,) if draft.report_ref else (draft.chart_no, draft.exam_date, draft.procedure_type)
            if key in seen:
                issues.append(Issue(sheet.sheet, row_no, "DUPLICATE_IN_FILE"))
                continue
            seen.add(key)
            sheet.rows.append(draft)
        return sheet

    def _row(self, kind, sheet, row_no, cells, today, issues) -> RowDraft | None:
        def bad(field_name: str, code: str = "INVALID_VALUE"):
            issues.append(Issue(sheet, row_no, code, FIELD_LABELS[field_name]))

        # 필수 값
        chart_no = _text(cells.get("chart_no"))
        name = _text(cells.get("patient_name"))
        try:
            exam_date = _date(cells.get("exam_date"))
        except ValueError:
            bad("exam_date", "INVALID_REQUIRED")
            return None
        for field_name, value in (("chart_no", chart_no), ("patient_name", name), ("exam_date", exam_date)):
            if not value:
                bad(field_name, "REQUIRED_MISSING")
                return None
        if not CHART_NO.match(chart_no):
            bad("chart_no", "INVALID_REQUIRED")
            return None
        if len(name) > 30:
            bad("patient_name", "INVALID_REQUIRED")
            return None
        draft = RowDraft(kind=kind, sheet=sheet, row=row_no, chart_no=chart_no, patient_name=name, exam_date=exam_date)

        # 선택 값: 형식이 틀리면 비워서 가져오고 검토 목록에 남긴다.
        for field_name in ("entrusted_on", "reported_on"):
            try:
                setattr(draft, field_name, _date(cells.get(field_name)))
            except ValueError:
                bad(field_name)
        try:
            draft.age = _int(cells.get("age"), 0, 130)
        except ValueError:
            bad("age")
        try:
            draft.specimen_count = _int(cells.get("specimen_count"), 1, 99)
        except ValueError:
            bad("specimen_count")
        sex = _text(cells.get("sex"))
        if sex:
            draft.sex = self.config.sexes.get(sex.upper()) or self.config.sexes.get(sex)
            if draft.sex is None:
                bad("sex")
        procedure = _text(cells.get("procedure_type")).casefold()
        if procedure:
            labels = {label: code for code, label in self.config.procedure_types.items()}
            draft.procedure_type = procedure if procedure in self.config.procedure_types else labels.get(procedure)
            if draft.procedure_type is None:
                bad("procedure_type")
        if kind == "CLO" and draft.procedure_type is None:
            draft.procedure_type = "stomach"
        ref = _text(cells.get("report_ref"))
        if ref and not REPORT_REF.match(ref):
            bad("report_ref")
        else:
            draft.report_ref = ref
        for field_name, limit in (("site", 80), ("registered_by", 20), ("handed_by", 20), ("received_by", 20)):
            value = _text(cells.get(field_name))
            if len(value) > limit:
                bad(field_name)
            else:
                setattr(draft, field_name, value)

        # 결과: 조직검사는 등록 코드 표기, CLO·Stool은 짧은 판정만(A11). 문장은 저장하지 않는다.
        result = _text(cells.get("result"))
        if result:
            if self.config.kinds[kind]["result_format"] == "codes":
                try:
                    draft.findings = self.codebook.parse_findings(result)
                except NotationError:
                    draft.review_codes.append("RESULT_UNPARSED")
            elif self.config.short_result_pattern.match(result):
                draft.result_code = result
            else:
                draft.review_codes.append("RESULT_UNPARSED")
        elif draft.reported_on:
            draft.review_codes.append("RESULT_UNPARSED")
        sydney = _text(cells.get("sydney"))
        if sydney and draft.procedure_type == "stomach":
            try:
                draft.sydney = self.codebook.parse_sydney(sydney)
            except NotationError:
                draft.review_codes.append("SYDNEY_UNPARSED")

        # CLO 대장처럼 회신날짜 열이 없는 대장은 결과가 적힌 행의 결과 날짜를 검사일로 본다(A18).
        if (draft.reported_on is None and (draft.result_code or draft.findings)
                and self.config.kinds[kind].get("result_date_default") == "exam_date"):
            draft.reported_on = draft.exam_date
        # 과거 행은 이미 인계된 것으로 본다. 오늘 행은 위탁일이 없으면 인계 대기로 남긴다.
        if (self.config.kinds[kind]["handoff"] and draft.entrusted_on is None
                and (draft.reported_on or draft.exam_date < today)):
            draft.entrusted_on = draft.exam_date
        return draft

    def _check_names_in_file(self, parsed: ParsedWorkbook) -> None:
        names: dict[str, set[str]] = {}
        for sheet in parsed.sheets:
            for draft in sheet.rows:
                names.setdefault(draft.chart_no, set()).add(draft.patient_name)
        conflicted = {chart for chart, values in names.items() if len(values) > 1}
        for sheet in parsed.sheets:
            for draft in [d for d in sheet.rows if d.chart_no in conflicted]:
                parsed.issues.append(Issue(sheet.sheet, draft.row, "NAME_MISMATCH", "수진자명"))
            sheet.rows = [d for d in sheet.rows if d.chart_no not in conflicted]


# ---------- 출력

LEDGER_HEADERS = ("검사일", "접수번호", "차트번호", "수진자명", "나이", "성별", "위탁일", "접수인", "검사종류",
                  "검사위치", "검체수량", "인계인", "인수인", "결과보고일", "조직검사결과", "특이소견", "비고란")
SHORT_HEADERS = ("검사일", "접수번호", "차트번호", "수진자명", "나이", "성별", "위탁일", "접수인",
                 "인계인", "인수인", "결과보고일", "검사결과")


def _status(entry: LedgerEntry, stage: str) -> str:
    if entry.followup_done_on:
        return f"재검 완료 {entry.followup_done_on.isoformat()}"
    if stage == "CLOSED":
        return "재검 불필요"
    return {"FOLLOWUP": "재검 예정", "DOCTOR_PENDING": "원장 지시 대기",
            "RESULT_PENDING": "결과 대기", "HANDOFF_PENDING": "인계 대기"}[stage]


def _header(sheet, headers) -> None:
    sheet.append(headers)
    for index, _ in enumerate(headers, 1):
        cell = sheet.cell(1, index)
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = 14


def _append(sheet, values) -> None:
    sheet.append([None if value == "" else value for value in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            # '='로 시작하는 값이 수식으로 실행되지 않게 문자열로 고정한다.
            cell.data_type = "s"
        elif isinstance(cell.value, date):
            cell.number_format = "yyyy.mm.dd"


# 대장별 출력 열(config kinds.*.export_headers)에 쓸 수 있는 열 이름과 값.
EXPORT_COLUMNS = {
    "검사일": lambda e: e.exam_date, "접수번호": lambda e: e.report_ref, "차트번호": lambda e: e.chart_no,
    "수진자명": lambda e: e.patient_name, "나이": lambda e: e.age, "성별": lambda e: e.sex,
    "CLO결과": lambda e: e.result_notation, "검사결과": lambda e: e.result_notation,
    "결과보고일": lambda e: e.reported_on, "위탁일": lambda e: e.entrusted_on,
}


def build_workbook(rows: list[tuple[LedgerEntry, str]], config: LedgerAppConfig, codebook: ResultCodeBook,
                   kinds: tuple[str, ...] = KINDS) -> Workbook:
    """DB 사본 엑셀. 대장별 시트. 수기 엑셀과 같은 열이 정해진 대장(CLO 5열, A18)은 그 열만, 나머지는
    공통 양식(조직검사 17열, 그 밖 12열) + 재검 2열. 이 파일을 고쳐 DB로 되돌리지 않는다."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    followup_headers = tuple(config.export["followup_headers"])
    for kind in kinds:
        sheet = workbook.create_sheet(config.kinds[kind]["sheet_names"][0])
        fixed = config.kinds[kind].get("export_headers")
        if fixed:
            _header(sheet, tuple(fixed))
            sheet.freeze_panes = "A2"
            for entry, _ in sorted((r for r in rows if r[0].kind == kind), key=lambda r: (r[0].exam_date, r[0].chart_no)):
                _append(sheet, [EXPORT_COLUMNS[name](entry) for name in fixed])
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(fixed))}{max(sheet.max_row, 1)}"
            continue
        headers = (LEDGER_HEADERS if kind == "BIOPSY" else SHORT_HEADERS) + followup_headers
        _header(sheet, headers)
        sheet.freeze_panes = "E2"
        for entry, stage in sorted((r for r in rows if r[0].kind == kind), key=lambda r: (r[0].exam_date, r[0].chart_no)):
            common = [entry.exam_date, entry.report_ref, entry.chart_no, entry.patient_name, entry.age, entry.sex,
                      entry.entrusted_on, entry.registered_by]
            if kind == "BIOPSY":
                values = common + [entry.procedure_type, entry.site, entry.specimen_count, entry.handed_by,
                                   entry.received_by, entry.reported_on, entry.result_notation,
                                   codebook.format_sydney(dict(entry.sydney)), ""]
            else:
                values = common + [entry.handed_by, entry.received_by, entry.reported_on, entry.result_notation]
            _append(sheet, values + [entry.followup_due, _status(entry, stage)])
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(sheet.max_row, 1)}"
    return workbook


FOLLOWUP_HEADERS = ("재검 예정일", "상태", "장부", "검사종류", "검사일", "차트번호", "수진자명", "나이", "성별",
                    "결과", "원장 지시일")


def build_followup_workbook(rows: list[tuple[LedgerEntry, str]], title: str, kind_labels: dict[str, str]) -> Workbook:
    """재검 대상자 목록. 화면에서 고른 조건(도래·기간·종류)의 행만, 재검일 순서로 담는다."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "재검 대상"
    sheet.append([title])
    sheet["A1"].font = Font(bold=True)
    sheet.append([])
    sheet.append(FOLLOWUP_HEADERS)
    for index, _ in enumerate(FOLLOWUP_HEADERS, 1):
        cell = sheet.cell(3, index)
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = 13
    sheet.freeze_panes = "A4"
    for entry, status in rows:
        _append(sheet, [entry.followup_due, status, kind_labels[entry.kind], entry.procedure_label, entry.exam_date,
                        entry.chart_no, entry.patient_name, entry.age, entry.sex, entry.result_notation,
                        entry.decided_at.astimezone(KST).date() if entry.decided_at else None])
    sheet.auto_filter.ref = f"A3:{get_column_letter(len(FOLLOWUP_HEADERS))}{max(sheet.max_row, 3)}"
    return workbook


# ---------- 종이 대장 양식 (A15). 열 이름·순서는 원내 종이 대장과 같다.
PAPER_FORMS = {
    "specimen": {
        "title": "검체 관리 장부", "short": "① 검체 관리 장부", "file_prefix": "검체관리장부",
        "headers": ("시행날짜", "차트번호", "성명", "검사명\n(위/대장내시경)", "조직검사 유무, 위치 및 개수",
                    "RUT\n유무", "RUT\n결과"),
        "keys": ("exam_date", "chart_no", "name", "procedure", "biopsy_text", "rut_done", "rut_result"),
        "widths": (12, 12, 12, 16, 44, 8, 10),
    },
    "pathology": {
        "title": "조직병리검체 접수, 인계, 인수 및 회신결과 대장 (위/대장)", "short": "② 조직병리검체 대장",
        "file_prefix": "조직병리검체대장",
        "headers": ("날짜", "접수번호", "차트번호", "성명", "성별/나이", "부위", "접수인(서명)", "위탁날짜",
                    "인계인(서명)", "수령인(서명)", "회신날짜", "결과"),
        "keys": ("exam_date", "report_ref", "chart_no", "name", "sex_age", "site", "registered_by", "entrusted_on",
                 "handed_by", "received_by", "reported_on", "result"),
        "widths": (11, 11, 10, 12, 9, 14, 11, 11, 11, 11, 11, 28),
    },
}


def format_sex_age(sex: str | None, age: int | None) -> str:
    """종이 대장 표기 M/65. 빈 쪽은 비워 둔다."""
    if not sex and age is None:
        return ""
    return f"{sex or ''}/{age if age is not None else ''}"


def build_paper_workbook(form: str, rows: list[dict], start: date, end: date, printed: datetime) -> Workbook:
    """종이 대장과 같은 열 순서·제목으로 출력한다. A4 가로, 열 이름 줄은 쪽마다 반복."""
    spec = PAPER_FORMS[form]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = spec["file_prefix"]
    columns = len(spec["headers"])
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    sheet["A1"] = spec["title"]
    sheet["A1"].font = Font(bold=True, size=16, underline="single")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    sheet["A2"] = (f"기간 {start:%Y.%m.%d} ~ {end:%Y.%m.%d} · {len(rows)}줄 · 출력 {printed:%Y.%m.%d %H:%M}"
                   " · CoriOps DB 사본(종이 대장 대조용)")
    sheet["A2"].font = Font(size=9, color="555555")
    sheet["A2"].alignment = Alignment(horizontal="right")
    for index, (label, width) in enumerate(zip(spec["headers"], spec["widths"]), 1):
        cell = sheet.cell(3, index, label)
        cell.fill = PatternFill("solid", fgColor="EEEAE0")
        cell.font = Font(bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[3].height = 30
    for row in rows:
        _append(sheet, [row[key] if row[key] is not None else "" for key in spec["keys"]])
        sheet.row_dimensions[sheet.max_row].height = 22
        for cell in sheet[sheet.max_row]:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A4"
    sheet.print_title_rows = "3:3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    return workbook
