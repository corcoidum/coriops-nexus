"""합성 자료를 장부로 옮기는 기능. 실제 수집·임상 해석은 포함하지 않는다."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date
from hashlib import sha256
from importlib.resources import files
from io import BytesIO
import json
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import yaml


type RecordKey = tuple[str, str]
EXCEL_CELL_TEXT_LIMIT = 32767


@dataclass(frozen=True, repr=False)
class ResultRecord:
    """출처 의미를 확인한 입력. 원문은 메모리 비교에만 사용한다."""

    report_ref: str
    chart_no: str
    patient_name: str
    exam_date: date | None = None
    age: str | None = None
    sex: str | None = None
    entrusted_date: date | None = None
    exam_kind: str | None = None
    specimen_count: int | None = None
    report_date: date | None = None
    site_text: str | None = None
    result_text: str | None = None


@dataclass(frozen=True)
class Column:
    header: str
    kind: str
    width: int
    field: str | None = None


@dataclass(frozen=True)
class LedgerPolicy:
    sheet_name: str
    freeze_panes: str
    review_label: str
    columns: tuple[Column, ...]

    @property
    def headers(self) -> tuple[str, ...]:
        return tuple(column.header for column in self.columns)


@dataclass(frozen=True)
class ReviewIssue:
    """개인값을 담지 않는 검토 위치·사유."""

    row: int
    column: str
    code: str


@dataclass(frozen=True, repr=False)
class MergeResult:
    appended: int
    unchanged: int
    issues: tuple[ReviewIssue, ...]
    fingerprints: dict[RecordKey, str]


def load_policy(path: Path | None = None) -> LedgerPolicy:
    """실제 Import 매핑과 분리된 합성 출력 정책을 읽는다."""
    # 설치된 패키지에서도 같은 정책을 읽도록 package resource로 둔다.
    source = Path(path) if path else files(__package__).joinpath("ledger_mapping.yaml")
    data = yaml.safe_load(source.read_text(encoding="utf-8"))
    columns = tuple(Column(**column) for column in data["columns"])
    allowed_fields = set(ResultRecord.__dataclass_fields__) - {"site_text", "result_text"}
    if (
        len(columns) != 17
        or len({column.header for column in columns}) != 17
        or any(column.kind not in {"date", "text", "integer", "manual", "abbreviation_pending"} for column in columns)
        or any(column.field not in allowed_fields for column in columns if column.field)
        or any(column.field for column in columns if column.kind in {"manual", "abbreviation_pending"})
    ):
        raise ValueError("INVALID_LEDGER_POLICY")
    return LedgerPolicy(data["sheet_name"], data["freeze_panes"], data["review_label"], columns)


def new_workbook(policy: LedgerPolicy) -> Workbook:
    """요청된 17열만 있는 시연용 workbook을 만든다."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = policy.sheet_name
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = policy.freeze_panes
    sheet.append(policy.headers)
    sheet.row_dimensions[1].height = 32
    for index, column in enumerate(policy.columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = column.width
        cell = sheet.cell(1, index)
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return workbook


def _validate_record(record: ResultRecord) -> RecordKey:
    for value in (record.report_ref, record.chart_no, record.patient_name):
        if not isinstance(value, str) or not value.strip():
            raise ValueError("MISSING_RECORD_KEY_OR_NAME")
    for value in (record.exam_date, record.entrusted_date, record.report_date):
        if value is not None and type(value) is not date:
            raise ValueError("INVALID_DATE_TYPE")
    for value in (record.age, record.sex, record.exam_kind, record.site_text, record.result_text):
        if value is not None and not isinstance(value, str):
            raise ValueError("INVALID_TEXT_TYPE")
    # Excel의 자동 잘림·XML 직렬화 오류가 발생하기 전에 배치 전체를 거부한다.
    for value in (record.report_ref, record.chart_no, record.patient_name, record.age, record.sex, record.exam_kind):
        if value is None:
            continue
        if len(value) > EXCEL_CELL_TEXT_LIMIT:
            raise ValueError("EXCEL_TEXT_TOO_LONG")
        if ILLEGAL_CHARACTERS_RE.search(value) or "\ufffe" in value or "\uffff" in value:
            raise ValueError("INVALID_EXCEL_TEXT")
        try:
            value.encode("utf-8")
        except UnicodeError:
            raise ValueError("INVALID_EXCEL_TEXT") from None
    if record.specimen_count is not None and (
        type(record.specimen_count) is not int or record.specimen_count < 1
    ):
        raise ValueError("INVALID_SPECIMEN_COUNT")
    return record.report_ref, record.chart_no


def merge_records(
    workbook: Workbook,
    records: Iterable[ResultRecord],
    policy: LedgerPolicy,
    fingerprints: Mapping[RecordKey, str] | None = None,
) -> MergeResult:
    """기존 행은 보존하고 신규 합성 행만 추가한다. 변경은 검토에 남긴다."""
    sheet = workbook[policy.sheet_name]
    if tuple(cell.value for cell in sheet[1]) != policy.headers:
        raise ValueError("HEADER_MISMATCH")
    report_column = next(i for i, column in enumerate(policy.columns, 1) if column.field == "report_ref")
    chart_column = next(i for i, column in enumerate(policy.columns, 1) if column.field == "chart_no")
    existing: dict[RecordKey, int] = {}
    for row in range(2, sheet.max_row + 1):
        values = tuple(sheet.cell(row, column).value for column in range(1, len(policy.columns) + 1))
        if all(value is None for value in values):
            continue
        key = (sheet.cell(row, report_column).value, sheet.cell(row, chart_column).value)
        if any(value is None or (isinstance(value, str) and not value.strip()) for value in key):
            raise ValueError("EXISTING_ROW_KEY_MISSING")
        # 숫자 셀은 앞자리 0이 사라져 문자열 키와 같은지 알 수 없으므로 추측해 맞추지 않는다.
        if any(not isinstance(value, str) for value in key):
            raise ValueError("EXISTING_ROW_KEY_NOT_TEXT")
        if key in existing:
            raise ValueError("DUPLICATE_EXISTING_KEY")
        existing[key] = row

    # 배치 입력 전체를 검증한 뒤 수정하여 뒤쪽 오류로 일부만 추가되지 않게 한다.
    batch = tuple(records)
    keys = tuple(_validate_record(record) for record in batch)
    known = {key: value for key, value in (fingerprints or {}).items() if key in existing}
    issues: list[ReviewIssue] = []
    appended = unchanged = 0
    for record, key in zip(batch, keys, strict=True):
        payload = json.dumps(asdict(record), ensure_ascii=True, sort_keys=True, default=str)
        digest = sha256(payload.encode("utf-8")).hexdigest()
        if key in existing:
            if key not in known:
                issues.append(ReviewIssue(existing[key], "", "EXISTING_ROW_REVIEW"))
            elif known[key] != digest:
                issues.append(ReviewIssue(existing[key], "", "SOURCE_CHANGED"))
            else:
                unchanged += 1
            continue

        row = sheet.max_row + 1
        for index, column in enumerate(policy.columns, 1):
            value = getattr(record, column.field) if column.field else None
            if column.kind == "abbreviation_pending":
                value = policy.review_label
                issues.append(ReviewIssue(row, column.header, "ABBREVIATION_PENDING"))
            cell = sheet.cell(row, index, value)
            cell.font = Font(name="Arial", size=11, color="243746")
            cell.alignment = Alignment(vertical="center", horizontal="right" if column.kind == "integer" else "left")
            if isinstance(value, str):
                # 외부 문자열이 = 등으로 시작해도 수식으로 실행되지 않게 한다.
                cell.data_type = "s"
                cell.number_format = "@"
            elif column.kind == "date":
                cell.number_format = "yyyy-mm-dd"
            elif column.kind == "integer":
                cell.number_format = "0"
            if column.kind == "abbreviation_pending":
                sheet.conditional_formatting.add(
                    cell.coordinate,
                    CellIsRule(operator="equal", formula=[f'"{policy.review_label}"'], fill=PatternFill("solid", fgColor="FFF0CC")),
                )
        sheet.row_dimensions[row].height = 26
        existing[key] = row
        known[key] = digest
        appended += 1
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(policy.columns))}{sheet.max_row}"
    return MergeResult(appended, unchanged, tuple(issues), known)


def save_new_workbook(workbook: Workbook, path: Path) -> Path:
    """새 파일에만 저장한다. 검증·직렬화 실패와 기존 파일 덮어쓰기를 구분한다."""
    path = Path(path)
    content = BytesIO()
    workbook.save(content)
    path.parent.mkdir(parents=True, exist_ok=True)
    # xb는 기존 파일·심볼릭 링크가 있으면 실패하므로 원본을 덮어쓰지 않는다.
    with path.open("xb") as destination:
        destination.write(content.getvalue())
    return path
